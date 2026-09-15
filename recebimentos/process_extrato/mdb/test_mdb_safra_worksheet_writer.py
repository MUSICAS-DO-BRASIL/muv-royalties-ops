from datetime import date
from decimal import Decimal

from openpyxl import Workbook, load_workbook
import pytest

from bank_extraction_core import SafraWorksheetRow
from mdb_safra_worksheet_writer import MdbSafraWorksheetWriter


HEADERS = ["data", "lancamento", "complemento", "documento", "valor_str", "valor", "Fonte Pagadora"]


def _book(path, headers=HEADERS, *, safra=True):
    book = Workbook()
    sheet = book.active
    sheet.title = "Safra" if safra else "Outra"
    sheet.append(headers)
    sentinel = book.create_sheet("Sentinela")
    sentinel["A1"] = "preservar"
    book.save(path)


def _rows(status="PASS"):
    return (
        SafraWorksheetRow(date(2026, 9, 1), "Lançamento A", "complemento multiline", "DOC-1", "1.234,56", Decimal("1234.56"), "Fonte A", status),
        SafraWorksheetRow(date(2026, 9, 2), "Lançamento B", None, None, "10,00", Decimal("10.00"), "Fonte B", status),
        SafraWorksheetRow(date(2026, 9, 3), "Lançamento C", "detalhe", "DOC-3", "0,01", Decimal("0.01"), "Fonte C", status),
    )


def test_writer_writes_complete_a_to_g_and_preserves_unrelated_sheet(tmp_path):
    path = tmp_path / "safra-sintetico.xlsx"
    _book(path)
    MdbSafraWorksheetWriter().write(path, _rows())

    book = load_workbook(path, data_only=False)
    sheet = book["Safra"]
    assert sheet["A2"].value.date() == date(2026, 9, 1)
    assert tuple(cell.value for cell in sheet[2][1:4]) == ("Lançamento A", "complemento multiline", "DOC-1")
    assert sheet["E2"].value == "1.234,56"
    assert Decimal(str(sheet["F2"].value)) == Decimal("1234.56")
    assert sheet["G2"].value == "Fonte A"
    assert (sheet["A3"].value.date(), *[cell.value for cell in sheet[3][1:4]], sheet["E3"].value, Decimal(str(sheet["F3"].value)), sheet["G3"].value) == (date(2026, 9, 2), "Lançamento B", None, None, "10,00", Decimal("10.00"), "Fonte B")
    assert [sheet.cell(row, 2).value for row in range(2, 5)] == ["Lançamento A", "Lançamento B", "Lançamento C"]
    assert book["Sentinela"]["A1"].value == "preservar"
    assert sheet.max_column == 7
    book.close()


def test_writer_accepts_operational_headers_linked_to_bs(tmp_path):
    path = tmp_path / "linked-headers.xlsx"
    _book(path, [f"=bs!{column}1" for column in "ABCDEF"] + ["Fonte Pagadora"])
    MdbSafraWorksheetWriter().write(path, _rows())
    book = load_workbook(path, data_only=False)
    assert [book["Safra"].cell(1, column).value for column in range(1, 7)] == [f"=bs!{column}1" for column in "ABCDEF"]
    assert book["Safra"]["G2"].value == "Fonte A"
    book.close()


@pytest.mark.parametrize("headers,safra", [
    (HEADERS[:-1], True),
    (["data", "lancamento", "complemento", "documento", "valor_str", "valor", "outra"], True),
    (HEADERS, False),
])
def test_writer_blocks_invalid_headers_or_missing_safra_sheet(tmp_path, headers, safra):
    path = tmp_path / "invalid.xlsx"
    _book(path, headers, safra=safra)
    with pytest.raises(ValueError):
        MdbSafraWorksheetWriter().write(path, _rows())


def test_writer_blocks_review_rows_and_existing_data(tmp_path):
    path = tmp_path / "review.xlsx"
    _book(path)
    with pytest.raises(ValueError):
        MdbSafraWorksheetWriter().write(path, _rows("REVIEW"))

    _book(path)
    book = load_workbook(path)
    book["Safra"]["A2"] = date(2026, 8, 1)
    book.save(path)
    book.close()
    with pytest.raises(ValueError):
        MdbSafraWorksheetWriter().write(path, _rows())
