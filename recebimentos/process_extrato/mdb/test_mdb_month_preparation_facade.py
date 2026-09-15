from decimal import Decimal
from pathlib import Path
import json
from hashlib import sha256
from dataclasses import replace

import pdfplumber
from openpyxl import Workbook, load_workbook

from bank_extraction_core import SafraAdapter, SafraExtractionResult
from mdb_month_preparation_facade import MdbMonthPreparationFacade


class _Page:
    def extract_text(self, **_):
        return """Extrato de Movimentação TESTE
CNPJ: 000
AG: 000 | CONTA: 000-000
Período de 01/09/2026 a 30/09/2026
LANÇAMENTOS REALIZADOS
01/09 TED RECEBIDA BCO 000 PAGADOR TESTE 0001 1.234,56
complemento sintético"""


class _Pdf:
    pages = [_Page()]
    def __enter__(self): return self
    def __exit__(self, *_): return False


def test_mdb_facade_happy_path_uses_real_safra_pipeline(monkeypatch, tmp_path):
    template = tmp_path / "template.xlsx"
    book = Workbook(); sheet = book.active; sheet.title = "Safra"
    sheet.append(["data", "lancamento", "complemento", "documento", "valor_str", "valor", "Fonte Pagadora"])
    book.create_sheet("Sentinela")["A1"] = "sintinela"
    book.save(template)
    source = tmp_path / "source.pdf"; source.write_bytes(b"synthetic")
    source_map = tmp_path / "map.json"; source_map.write_text(json.dumps({"aliases":[{"bank_payor_alias":"PAGADOR TESTE","canonical_royalty_source":"Fonte Teste"}]}), encoding="utf-8")
    monkeypatch.setenv("MUV_SAFRA_SOURCE_MAP", str(source_map))
    monkeypatch.setattr(pdfplumber, "open", lambda _: _Pdf())
    result = SafraAdapter().extract(source, "2026-09")
    assert isinstance(result, SafraExtractionResult)
    outcome = MdbMonthPreparationFacade(template_path=template, monthly_root=tmp_path).prepare(result, "2026-09")
    expected = tmp_path / "2026" / "092026" / "MDB" / "Conciliação - Músicas do Brasil_202609.xlsx"
    assert outcome.status == "PREPARED" and outcome.workbook_path == expected and expected.is_file()
    output = load_workbook(expected); safra = output["Safra"]; row = result.safra_rows[0]
    assert safra["A2"].value.date() == row.transaction_date
    assert [safra.cell(2, c).value for c in range(2, 6)] == [row.lancamento, row.complemento, row.documento, row.valor_str]
    assert Decimal(str(safra["F2"].value)) == row.credit and safra["G2"].value == row.payor_source
    assert output["Sentinela"]["A1"].value == "sintinela"
    output.close()


def test_mdb_facade_second_run_is_create_only(monkeypatch, tmp_path):
    template = tmp_path / "template.xlsx"
    book = Workbook(); sheet = book.active; sheet.title = "Safra"
    sheet.append(["data", "lancamento", "complemento", "documento", "valor_str", "valor", "Fonte Pagadora"])
    book.create_sheet("Sentinela")["A1"] = "sintinela"; book.save(template)
    template_hash = sha256(template.read_bytes()).hexdigest()
    source = tmp_path / "source.pdf"; source.write_bytes(b"synthetic")
    source_map = tmp_path / "map.json"; source_map.write_text(json.dumps({"aliases":[{"bank_payor_alias":"PAGADOR TESTE","canonical_royalty_source":"Fonte Teste"}]}), encoding="utf-8")
    monkeypatch.setenv("MUV_SAFRA_SOURCE_MAP", str(source_map)); monkeypatch.setattr(pdfplumber, "open", lambda _: _Pdf())
    result = SafraAdapter().extract(source, "2026-09")
    facade = MdbMonthPreparationFacade(template_path=template, monthly_root=tmp_path)
    first = facade.prepare(result, "2026-09")
    assert first.status == "PREPARED" and first.workbook_path is not None
    first_hash = sha256(first.workbook_path.read_bytes()).hexdigest(); first_size = first.workbook_path.stat().st_size
    first_rows = load_workbook(first.workbook_path)["Safra"].max_row
    second = facade.prepare(result, "2026-09")
    assert second.status == "ALREADY_EXISTS" and second.workbook_path == first.workbook_path
    assert sha256(first.workbook_path.read_bytes()).hexdigest() == first_hash
    assert first.workbook_path.stat().st_size == first_size
    assert load_workbook(first.workbook_path)["Safra"].max_row == first_rows
    assert sha256(template.read_bytes()).hexdigest() == template_hash


def test_mdb_facade_writer_failure_does_not_publish_final_file(monkeypatch, tmp_path):
    template = tmp_path / "template.xlsx"
    book = Workbook(); sheet = book.active; sheet.title = "Safra"
    sheet.append(["data", "lancamento", "complemento", "documento", "valor_str", "valor", "Fonte Pagadora"]); book.save(template)
    template_hash = sha256(template.read_bytes()).hexdigest()
    source = tmp_path / "source.pdf"; source.write_bytes(b"synthetic")
    source_map = tmp_path / "map.json"; source_map.write_text(json.dumps({"aliases":[{"bank_payor_alias":"PAGADOR TESTE","canonical_royalty_source":"Fonte Teste"}]}))
    monkeypatch.setenv("MUV_SAFRA_SOURCE_MAP", str(source_map)); monkeypatch.setattr(pdfplumber, "open", lambda _: _Pdf())
    result = SafraAdapter().extract(source, "2026-09")
    class FailingWriter:
        def write(self, *_args, **_kwargs): raise RuntimeError("synthetic writer failure")
    outcome = MdbMonthPreparationFacade(template_path=template, monthly_root=tmp_path, writer=FailingWriter()).prepare(result, "2026-09")
    final = tmp_path / "2026" / "092026" / "MDB" / "Conciliação - Músicas do Brasil_202609.xlsx"
    assert outcome.status == "BLOCKED" and not final.exists()
    assert sha256(template.read_bytes()).hexdigest() == template_hash


def test_mdb_facade_blocks_review_rich_row(monkeypatch, tmp_path):
    template = tmp_path / "template.xlsx"
    book = Workbook(); sheet = book.active; sheet.title = "Safra"
    sheet.append(["data", "lancamento", "complemento", "documento", "valor_str", "valor", "Fonte Pagadora"]); book.save(template)
    template_hash = sha256(template.read_bytes()).hexdigest()
    source = tmp_path / "source.pdf"; source.write_bytes(b"synthetic")
    source_map = tmp_path / "map.json"; source_map.write_text(json.dumps({"aliases":[{"bank_payor_alias":"PAGADOR TESTE","canonical_royalty_source":"Fonte Teste"}]}))
    monkeypatch.setenv("MUV_SAFRA_SOURCE_MAP", str(source_map)); monkeypatch.setattr(pdfplumber, "open", lambda _: _Pdf())
    result = SafraAdapter().extract(source, "2026-09")
    review_row = replace(result.safra_rows[0], status="REVIEW")
    reviewed = replace(result, safra_rows=(review_row,))
    outcome = MdbMonthPreparationFacade(template_path=template, monthly_root=tmp_path).prepare(reviewed, "2026-09")
    final = tmp_path / "2026" / "092026" / "MDB" / "Conciliação - Músicas do Brasil_202609.xlsx"
    assert outcome.status == "BLOCKED" and not final.exists()
    assert sha256(template.read_bytes()).hexdigest() == template_hash


def test_mdb_facade_blocks_missing_template(monkeypatch, tmp_path):
    source = tmp_path / "source.pdf"; source.write_bytes(b"synthetic")
    source_map = tmp_path / "map.json"; source_map.write_text(json.dumps({"aliases":[{"bank_payor_alias":"PAGADOR TESTE","canonical_royalty_source":"Fonte Teste"}]}))
    monkeypatch.setenv("MUV_SAFRA_SOURCE_MAP", str(source_map)); monkeypatch.setattr(pdfplumber, "open", lambda _: _Pdf())
    result = SafraAdapter().extract(source, "2026-09")
    outcome = MdbMonthPreparationFacade(template_path=tmp_path / "missing.xlsx", monthly_root=tmp_path).prepare(result, "2026-09")
    final = tmp_path / "2026" / "092026" / "MDB" / "Conciliação - Músicas do Brasil_202609.xlsx"
    assert outcome.status == "BLOCKED" and not final.exists()


def test_mdb_facade_blocks_invalid_template(monkeypatch, tmp_path):
    template = tmp_path / "invalid_template.xlsx"; template.write_text("synthetic invalid xlsx", encoding="utf-8")
    template_hash = sha256(template.read_bytes()).hexdigest()
    source = tmp_path / "source.pdf"; source.write_bytes(b"synthetic")
    source_map = tmp_path / "map.json"; source_map.write_text(json.dumps({"aliases":[{"bank_payor_alias":"PAGADOR TESTE","canonical_royalty_source":"Fonte Teste"}]}))
    monkeypatch.setenv("MUV_SAFRA_SOURCE_MAP", str(source_map)); monkeypatch.setattr(pdfplumber, "open", lambda _: _Pdf())
    result = SafraAdapter().extract(source, "2026-09")
    outcome = MdbMonthPreparationFacade(template_path=template, monthly_root=tmp_path).prepare(result, "2026-09")
    final = tmp_path / "2026" / "092026" / "MDB" / "Conciliação - Músicas do Brasil_202609.xlsx"
    assert outcome.status == "BLOCKED" and not final.exists()
    assert sha256(template.read_bytes()).hexdigest() == template_hash


def test_mdb_facade_blocks_template_without_safra_sheet(monkeypatch, tmp_path):
    template = tmp_path / "no-safra.xlsx"
    book = Workbook(); book.active.title = "Sentinela"; book.save(template)
    template_hash = sha256(template.read_bytes()).hexdigest()
    source = tmp_path / "source.pdf"; source.write_bytes(b"synthetic")
    source_map = tmp_path / "map.json"; source_map.write_text(json.dumps({"aliases":[{"bank_payor_alias":"PAGADOR TESTE","canonical_royalty_source":"Fonte Teste"}]}))
    monkeypatch.setenv("MUV_SAFRA_SOURCE_MAP", str(source_map)); monkeypatch.setattr(pdfplumber, "open", lambda _: _Pdf())
    result = SafraAdapter().extract(source, "2026-09")
    outcome = MdbMonthPreparationFacade(template_path=template, monthly_root=tmp_path).prepare(result, "2026-09")
    final = tmp_path / "2026" / "092026" / "MDB" / "Conciliação - Músicas do Brasil_202609.xlsx"
    assert outcome.status == "BLOCKED" and not final.exists()
    assert sha256(template.read_bytes()).hexdigest() == template_hash


def test_mdb_facade_blocks_invalid_safra_header(monkeypatch, tmp_path):
    template = tmp_path / "invalid-header.xlsx"
    book = Workbook(); sheet = book.active; sheet.title = "Safra"
    sheet.append(["data", "lancamento", "complemento", "documento", "valor_str", "valor", "Fonte Incorreta"]); book.save(template)
    source = tmp_path / "source.pdf"; source.write_bytes(b"synthetic")
    source_map = tmp_path / "map.json"; source_map.write_text(json.dumps({"aliases":[{"bank_payor_alias":"PAGADOR TESTE","canonical_royalty_source":"Fonte Teste"}]}))
    monkeypatch.setenv("MUV_SAFRA_SOURCE_MAP", str(source_map)); monkeypatch.setattr(pdfplumber, "open", lambda _: _Pdf())
    result = SafraAdapter().extract(source, "2026-09")
    outcome = MdbMonthPreparationFacade(template_path=template, monthly_root=tmp_path).prepare(result, "2026-09")
    final = tmp_path / "2026" / "092026" / "MDB" / "Conciliação - Músicas do Brasil_202609.xlsx"
    assert outcome.status == "BLOCKED" and not final.exists()


def test_mdb_facade_blocks_incompatible_bank_result(monkeypatch, tmp_path):
    template = tmp_path / "template.xlsx"
    book = Workbook(); sheet = book.active; sheet.title = "Safra"
    sheet.append(["data", "lancamento", "complemento", "documento", "valor_str", "valor", "Fonte Pagadora"])
    book.create_sheet("Sentinela")["A1"] = "sintinela"
    book.save(template)
    template_hash = sha256(template.read_bytes()).hexdigest()
    source = tmp_path / "source.pdf"; source.write_bytes(b"synthetic")
    source_map = tmp_path / "map.json"; source_map.write_text(json.dumps({"aliases":[{"bank_payor_alias":"PAGADOR TESTE","canonical_royalty_source":"Fonte Teste"}]}))
    monkeypatch.setenv("MUV_SAFRA_SOURCE_MAP", str(source_map)); monkeypatch.setattr(pdfplumber, "open", lambda _: _Pdf())
    result = replace(SafraAdapter().extract(source, "2026-09"), entity="HM")
    outcome = MdbMonthPreparationFacade(template_path=template, monthly_root=tmp_path).prepare(result, "2026-09")
    final = tmp_path / "2026" / "092026" / "MDB" / "Conciliação - Músicas do Brasil_202609.xlsx"
    assert outcome.status == "BLOCKED" and outcome.workbook_path is None
    assert outcome.errors == ("Resultado Safra inválido para MDB.",)
    assert not final.exists()
    assert sha256(template.read_bytes()).hexdigest() == template_hash


def test_mdb_facade_blocks_competence_mismatch(monkeypatch, tmp_path):
    template = tmp_path / "template.xlsx"
    book = Workbook(); sheet = book.active; sheet.title = "Safra"
    sheet.append(["data", "lancamento", "complemento", "documento", "valor_str", "valor", "Fonte Pagadora"])
    book.create_sheet("Sentinela")["A1"] = "sintinela"
    book.save(template)
    template_hash = sha256(template.read_bytes()).hexdigest()
    source = tmp_path / "source.pdf"; source.write_bytes(b"synthetic")
    source_map = tmp_path / "map.json"; source_map.write_text(json.dumps({"aliases":[{"bank_payor_alias":"PAGADOR TESTE","canonical_royalty_source":"Fonte Teste"}]}))
    monkeypatch.setenv("MUV_SAFRA_SOURCE_MAP", str(source_map)); monkeypatch.setattr(pdfplumber, "open", lambda _: _Pdf())
    result = replace(SafraAdapter().extract(source, "2026-09"), period="2026-08")
    assert result.entity == "MDB" and result.bank_source == "SAFRA" and result.safra_rows
    outcome = MdbMonthPreparationFacade(template_path=template, monthly_root=tmp_path).prepare(result, "2026-09")
    august_final = tmp_path / "2026" / "082026" / "MDB" / "Conciliação - Músicas do Brasil_202608.xlsx"
    september_final = tmp_path / "2026" / "092026" / "MDB" / "Conciliação - Músicas do Brasil_202609.xlsx"
    assert outcome.status == "BLOCKED" and outcome.workbook_path is None
    assert outcome.errors == ("Resultado Safra inválido para MDB.",)
    assert not august_final.exists() and not september_final.exists()
    assert sha256(template.read_bytes()).hexdigest() == template_hash


def test_mdb_facade_blocks_bank_source_mismatch(monkeypatch, tmp_path):
    template = tmp_path / "template.xlsx"
    book = Workbook(); sheet = book.active; sheet.title = "Safra"
    sheet.append(["data", "lancamento", "complemento", "documento", "valor_str", "valor", "Fonte Pagadora"])
    book.create_sheet("Sentinela")["A1"] = "sintinela"
    book.save(template)
    template_hash = sha256(template.read_bytes()).hexdigest()
    source = tmp_path / "source.pdf"; source.write_bytes(b"synthetic")
    source_map = tmp_path / "map.json"; source_map.write_text(json.dumps({"aliases":[{"bank_payor_alias":"PAGADOR TESTE","canonical_royalty_source":"Fonte Teste"}]}))
    monkeypatch.setenv("MUV_SAFRA_SOURCE_MAP", str(source_map)); monkeypatch.setattr(pdfplumber, "open", lambda _: _Pdf())
    result = replace(SafraAdapter().extract(source, "2026-09"), bank_source="BTG")
    assert result.entity == "MDB" and result.period == "2026-09" and result.safra_rows
    outcome = MdbMonthPreparationFacade(template_path=template, monthly_root=tmp_path).prepare(result, "2026-09")
    final = tmp_path / "2026" / "092026" / "MDB" / "Conciliação - Músicas do Brasil_202609.xlsx"
    assert outcome.status == "BLOCKED" and outcome.workbook_path is None
    assert outcome.errors == ("Resultado Safra inválido para MDB.",)
    assert not final.exists()
    assert sha256(template.read_bytes()).hexdigest() == template_hash
