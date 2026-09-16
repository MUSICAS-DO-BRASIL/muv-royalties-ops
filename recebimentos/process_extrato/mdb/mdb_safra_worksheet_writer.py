"""Typed, fail-closed writer for the MDB Safra A:G worksheet contract."""
from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime, time
import sys
from pathlib import Path
import unicodedata

from openpyxl import load_workbook

from bank_extraction_core import SafraWorksheetRow


EXPECTED_HEADERS = ("data", "lancamento", "complemento", "documento", "valor_str", "valor", "fonte pagadora")
LINKED_HEADER_FORMULAS = tuple(f"=bs!{column}1".casefold() for column in "ABCDEF")


@dataclass(frozen=True)
class MdbWorkbookLayout:
    safra_sheet_name: str


class UnsupportedWorkbookPlatformError(RuntimeError):
    """The operational workbook backend cannot run on this platform."""


class ExcelComWorkbookBackend:
    """Windows-only operational XLSX mutation backend; preserves native features."""
    def write_safra_rows(self, path: str | Path, rows: tuple[SafraWorksheetRow, ...], layout: MdbWorkbookLayout) -> None:
        if sys.platform != "win32":
            raise UnsupportedWorkbookPlatformError("Excel COM operational workbook backend requires Windows.")
        import pythoncom
        import win32com.client
        app = workbook = None
        try:
            pythoncom.CoInitialize()
            app = win32com.client.DispatchEx("Excel.Application")
            app.Visible = False; app.DisplayAlerts = False
            workbook = app.Workbooks.Open(str(Path(path).resolve()), 0, False)
            sheet = workbook.Worksheets.Item(layout.safra_sheet_name)
            for number, row in enumerate(rows, start=2):
                sheet.Cells(number, 1).Value = datetime.combine(row.transaction_date, time.min)
                sheet.Cells(number, 2).Value = row.lancamento
                sheet.Cells(number, 3).Value = row.complemento
                sheet.Cells(number, 4).Value = row.documento
                sheet.Cells(number, 5).Value = row.valor_str
                sheet.Cells(number, 6).Value = float(row.credit)
                sheet.Cells(number, 7).Value = row.payor_source
            app.CalculateFullRebuild()
            workbook.Save()
        finally:
            if workbook is not None: workbook.Close(False)
            if app is not None: app.Quit()
            pythoncom.CoUninitialize()


class SyntheticOpenpyxlWorkbookBackend:
    """Non-production backend; synthetic callers must inject it explicitly."""
    def write_safra_rows(self, path, rows, layout) -> None:
        workbook = load_workbook(path)
        try:
            sheet = workbook[layout.safra_sheet_name]
            for number, row in enumerate(rows, start=2):
                for column, value in enumerate((row.transaction_date, row.lancamento, row.complemento, row.documento, row.valor_str, row.credit, row.payor_source), start=1):
                    sheet.cell(number, column, value)
            workbook.save(path)
        finally:
            workbook.close()


def _fold(value: object) -> str:
    text = unicodedata.normalize("NFKD", str(value or ""))
    return "".join(character for character in text if not unicodedata.combining(character)).strip().casefold()


def _has_expected_headers(sheet) -> bool:
    headers = tuple(_fold(sheet.cell(1, column).value) for column in range(1, 8))
    if headers == EXPECTED_HEADERS:
        return True
    linked_headers = tuple(str(sheet.cell(1, column).value or "").strip().casefold() for column in range(1, 7))
    return linked_headers == LINKED_HEADER_FORMULAS and headers[6] == EXPECTED_HEADERS[6]


class MdbSafraWorksheetWriter:
    """Write validated Safra rich rows without touching MDB mappings or other sheets."""

    def __init__(self, backend: ExcelComWorkbookBackend | SyntheticOpenpyxlWorkbookBackend | None = None) -> None:
        self.backend = backend if backend is not None else ExcelComWorkbookBackend()

    def write(self, path: str | Path, rows: tuple[SafraWorksheetRow, ...], *, layout: MdbWorkbookLayout = MdbWorkbookLayout("Safra")) -> None:
        workbook = load_workbook(path, data_only=False, keep_links=True)
        try:
            if layout.safra_sheet_name not in workbook.sheetnames:
                raise ValueError("Aba Safra não encontrada.")
            sheet = workbook[layout.safra_sheet_name]
            if not _has_expected_headers(sheet):
                raise ValueError("Layout Safra MDB inválido: headers A:G não correspondem ao contrato.")
            if any(sheet.cell(row, column).value is not None for row in range(2, sheet.max_row + 1) for column in range(1, 8)):
                raise ValueError("Aba Safra MDB já contém dados; escrita não é idempotente.")
            if any(row.status != "PASS" for row in rows):
                raise ValueError("Linha Safra em revisão não pode ser publicada.")
        finally:
            workbook.close()
        self.backend.write_safra_rows(path, rows, layout)
