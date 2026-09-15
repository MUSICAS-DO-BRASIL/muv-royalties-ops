"""Read-only audit utilities for the HM 2026-08 structural staging copy."""

from __future__ import annotations

from dataclasses import dataclass, asdict
import json
from pathlib import Path
import re

from openpyxl import load_workbook


@dataclass(frozen=True)
class FormulaRecord:
    sheet: str
    cell: str
    formula: str


def formula_records(workbook_path: str | Path) -> dict[tuple[str, str], str]:
    workbook = load_workbook(Path(workbook_path), read_only=True, data_only=False, keep_links=True)
    return {
        (sheet.title, cell.coordinate): cell.value
        for sheet in workbook.worksheets
        for row in sheet.iter_rows()
        for cell in row
        if isinstance(cell.value, str) and cell.value.startswith("=")
    }


def discover_btg_dependencies(workbook_path: str | Path, bank_sheet: str) -> list[dict[str, str]]:
    """Find formula dependencies on the legacy BTG sheet without editing it."""
    reference = re.compile(r"(?:'" + re.escape(bank_sheet) + r"'|" + re.escape(bank_sheet) + r")!\$?([A-Z]+)")
    dependencies: list[dict[str, str]] = []
    for (sheet, cell), formula in formula_records(workbook_path).items():
        columns = reference.findall(formula)
        for column in columns:
            purpose = "bank_credit_by_source" if column == "C" else "bank_source_mapping" if column == "D" else "bank_reference"
            dependencies.append({
                "referencing_sheet": sheet,
                "cell": cell,
                "formula": formula,
                "referenced_btg_column": column,
                "semantic_purpose": purpose,
            })
    return dependencies


def formula_regression(source_path: str | Path, staging_path: str | Path, bank_sheet: str) -> list[dict[str, str]]:
    before = formula_records(source_path)
    after = formula_records(staging_path)
    keys = sorted(set(before) | set(after))
    report: list[dict[str, str]] = []
    for sheet, cell in keys:
        before_formula = before.get((sheet, cell), "")
        after_formula = after.get((sheet, cell), "")
        fallback_only = (
            sheet == bank_sheet and re.fullmatch(r"D(?:[2-9]|[1-9]\d{1,2})", cell) is not None
            and before_formula.replace('"xxxxERROExxxx"', '"REVIEW / UNKNOWN"') == after_formula
        )
        passed = before_formula == after_formula or fallback_only
        report.append({
            "cell": f"{sheet}!{cell}",
            "formula_before": before_formula,
            "formula_after": after_formula,
            "expected": "unchanged" if before_formula == after_formula else "same XLOOKUP with REVIEW / UNKNOWN fallback",
            "pass_fail": "PASS" if passed else "FAIL",
        })
    return report


def write_json(path: str | Path, value: object) -> Path:
    destination = Path(path)
    destination.parent.mkdir(parents=True, exist_ok=True)
    destination.write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding="utf-8")
    return destination
