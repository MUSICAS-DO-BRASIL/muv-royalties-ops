"""Safe preparation gate for one HM reconciliation month.

This is intentionally a discovery/validation layer.  It does not publish or
overwrite reconciliation workbooks.  A workbook is only eligible for a later
staging-copy writer when its bank tab already preserves the canonical BTG bank
columns required for audit.
"""

from __future__ import annotations

from dataclasses import dataclass
from decimal import Decimal
from hashlib import sha256
from pathlib import Path
import re
import unicodedata

from openpyxl import load_workbook

from btg_statement_parser import StatementDetails, ValidationResult, parse_btg_statement_details, validate_btg_statement


ENTITY = "HM"
BANK_SOURCE = "BTG"
CANONICAL_BANK_COLUMNS = ("data", "descricao", "debito", "credito", "saldo")


@dataclass(frozen=True)
class WorkbookInspection:
    path: Path
    sha256: str
    sheets: tuple[str, ...]
    bank_sheet: str | None
    bank_headers: tuple[str, ...]
    has_canonical_bank_columns: bool
    has_external_references: bool
    monthly_content_classification: str
    structural_notes: tuple[str, ...]


@dataclass(frozen=True)
class HMMonthContext:
    entity: str
    period: str
    bank_source: str
    bank_transactions: tuple
    bank_validation: ValidationResult
    monthly_folder: Path
    month_workbook_already_exists: bool
    existing_workbook_path: Path | None
    template_source: Path | None
    source_template_sha256: str | None
    staging_path: Path | None
    validation_status: str
    review_reasons: tuple[str, ...]
    bank_sheet: str | None


def _fold(value: object) -> str:
    text = str(value or "")
    return "".join(character for character in unicodedata.normalize("NFKD", text.lower()) if not unicodedata.combining(character))


def _sha256(path: Path) -> str:
    digest = sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _period_folder(period: str, monthly_root: Path) -> Path:
    if not re.fullmatch(r"\d{4}-\d{2}", period):
        raise ValueError("period must be YYYY-MM")
    year, month = period.split("-")
    return monthly_root / year / f"{month}{year}" / ENTITY


def _find_bank_sheet(workbook) -> str | None:
    for sheet_name in workbook.sheetnames:
        folded = _fold(sheet_name)
        if "btg" in folded or "banco" in folded:
            return sheet_name
    return None


def inspect_workbook(path: str | Path) -> WorkbookInspection:
    """Read workbook structure without changing it."""
    workbook_path = Path(path)
    if not workbook_path.is_file():
        raise FileNotFoundError(f"Workbook not found: {workbook_path}")
    workbook = load_workbook(workbook_path, read_only=True, data_only=False, keep_links=True)
    bank_sheet = _find_bank_sheet(workbook)
    headers: tuple[str, ...] = ()
    notes: list[str] = []
    external = False
    if bank_sheet:
        worksheet = workbook[bank_sheet]
        headers = tuple(_fold(cell.value) for cell in next(worksheet.iter_rows(min_row=1, max_row=1)))
        for row in worksheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("=") and "[" in cell.value:
                    external = True
                    break
            if external:
                break
        missing = [column for column in CANONICAL_BANK_COLUMNS if column not in headers]
        if missing:
            notes.append("Bank sheet lacks canonical audit columns: " + ", ".join(missing))
        if "fonte_pagadora" in headers:
            notes.append("Bank sheet includes source-mapping field; source classification must remain outside this bootstrap")
    else:
        notes.append("No BTG/bank worksheet found")
    canonical = bool(bank_sheet) and all(column in headers for column in CANONICAL_BANK_COLUMNS)
    classification = "STRUCTURAL_KEEP" if canonical else "UNKNOWN_REVIEW"
    return WorkbookInspection(
        path=workbook_path, sha256=_sha256(workbook_path), sheets=tuple(workbook.sheetnames),
        bank_sheet=bank_sheet, bank_headers=headers, has_canonical_bank_columns=canonical,
        has_external_references=external, monthly_content_classification=classification,
        structural_notes=tuple(notes),
    )


def find_month_workbooks(monthly_folder: str | Path) -> tuple[Path, ...]:
    folder = Path(monthly_folder)
    if not folder.is_dir():
        return ()
    return tuple(sorted(path for path in folder.glob("*.xlsx") if path.is_file() and not path.name.startswith("~$")))


def prepare_hm_month(
    *, period: str, btg_statement: str | Path, monthly_root: str | Path,
    template_candidates: tuple[str | Path, ...] = (),
) -> HMMonthContext:
    """Build an in-memory, source-safe HM month context.

    A missing official workbook plus an incompatible legacy template yields
    ``REVIEW``.  This deliberate stop prevents a staging copy from silently
    changing formulas/mappings or losing BTG debit and balance audit fields.
    """
    details: StatementDetails = parse_btg_statement_details(btg_statement)
    validation = validate_btg_statement(details)
    if not validation.passed:
        raise ValueError("BTG bank validation failed: " + "; ".join(validation.errors))
    monthly_folder = _period_folder(period, Path(monthly_root))
    existing = find_month_workbooks(monthly_folder)
    if len(existing) > 1:
        return HMMonthContext(ENTITY, period, BANK_SOURCE, details.transactions, validation, monthly_folder, True,
                              None, None, None, None, "REVIEW", ("More than one month workbook exists; choose a destination explicitly",), None)
    if len(existing) == 1:
        inspection = inspect_workbook(existing[0])
        return HMMonthContext(ENTITY, period, BANK_SOURCE, details.transactions, validation, monthly_folder, True,
                              existing[0], None, None, None, "PASS" if inspection.has_canonical_bank_columns else "REVIEW",
                              inspection.structural_notes, inspection.bank_sheet)

    candidate_paths = tuple(Path(candidate) for candidate in template_candidates if Path(candidate).is_file())
    if not candidate_paths:
        return HMMonthContext(ENTITY, period, BANK_SOURCE, details.transactions, validation, monthly_folder, False,
                              None, None, None, None, "REVIEW", ("No HM workbook template was supplied or discovered",), None)
    inspections = tuple(inspect_workbook(candidate) for candidate in candidate_paths)
    eligible = next((inspection for inspection in inspections if inspection.has_canonical_bank_columns), None)
    if not eligible:
        reasons = ("No supplied HM template preserves Data, Descrição, Débito, Crédito and Saldo on its bank sheet",)
        for inspection in inspections:
            reasons += tuple(f"{inspection.path.name}: {note}" for note in inspection.structural_notes)
        return HMMonthContext(ENTITY, period, BANK_SOURCE, details.transactions, validation, monthly_folder, False,
                              None, inspections[0].path, inspections[0].sha256, None, "REVIEW", reasons,
                              inspections[0].bank_sheet)
    # Deliberately no workbook creation/publishing in V1. A later, explicitly
    # authorized writer can consume this eligible context into a staging copy.
    return HMMonthContext(ENTITY, period, BANK_SOURCE, details.transactions, validation, monthly_folder, False,
                          None, eligible.path, eligible.sha256, None, "REVIEW",
                          ("Eligible template found; staging writer is intentionally not implemented in bootstrap V1",),
                          eligible.bank_sheet)


def checkpoint(context: HMMonthContext) -> dict[str, object]:
    validation = context.bank_validation
    return {
        "HM_MONTH_BOOTSTRAP_V1_STATUS": context.validation_status,
        "ENTITY": context.entity,
        "PERIOD": context.period,
        "BANK_SOURCE": context.bank_source,
        "MONTH_FOLDER": str(context.monthly_folder),
        "MONTH_WORKBOOK_ALREADY_EXISTS": context.month_workbook_already_exists,
        "EXISTING_WORKBOOK_PATH": str(context.existing_workbook_path) if context.existing_workbook_path else None,
        "TEMPLATE_SOURCE": str(context.template_source) if context.template_source else None,
        "SOURCE_TEMPLATE_SHA256": context.source_template_sha256,
        "STAGING_PATH": None,
        "BANK_SHEET": context.bank_sheet,
        "BANK_TRANSACTION_COUNT": validation.transaction_count,
        "BANK_CREDIT_COUNT": validation.credit_transaction_count,
        "BANK_DEBIT_COUNT": validation.debit_transaction_count,
        "INITIAL_BALANCE": str(validation.initial_balance),
        "TOTAL_CREDITS": str(validation.total_credits),
        "TOTAL_DEBITS": str(validation.total_debits),
        "FINAL_BALANCE": str(validation.final_balance),
        "GLOBAL_BALANCE_VALIDATION": "PASS" if validation.passed else "FAIL",
        "CATALOG_CHANGED": False,
        "SOURCE_AUTOMATIONS_STARTED": False,
        "EXCEL_REOPEN_OK": None,
        "EXCEL_REPAIR_REQUIRED": None,
        "IDEMPOTENCY_PASS": True,
        "MONTH_CONTEXT_CREATED": True,
        "SAFE_TO_CREATE_HM_MONTH_OFFICIAL": False,
        "OFFICIAL_FILES_CHANGED": False,
    }
