"""Safe, generic monthly reconciliation preparation for HM and MDB.

This service deliberately separates month discovery from creation.  It only
creates a new workbook from an explicitly clean project-side template, and it
uses create-only publication so an existing reconciliation can never be
overwritten.
"""
from __future__ import annotations

from dataclasses import dataclass
from datetime import date
from decimal import Decimal
from pathlib import Path
from tempfile import TemporaryDirectory
from typing import Callable, Literal
import shutil
import os
import sys

from openpyxl import load_workbook

from bank_extraction_core import BankExtractionResult, ENTITY_BANK, ROOT
from cloud_aware_workbook_publisher import create_workbook_only


TEMPLATE_ROOT = Path(os.environ.get("MUV_TEMPLATE_ROOT") or Path(__file__).resolve().parent / "templates").expanduser()
_NAMES = {
    "HM": "Conciliação - Hurst Music",
    "MDB": "Conciliação - Músicas do Brasil",
}


@dataclass(frozen=True)
class MonthPreparationResult:
    entity: str
    competence: str
    bank_source: str
    status: Literal["PREPARED", "REVIEW", "BLOCKED", "CONFLICT"]
    month_folder: Path
    workbook_path: Path | None
    already_existed: bool
    created_now: bool
    bank_rows: int
    bank_total: Decimal
    warnings: tuple[str, ...] = ()
    errors: tuple[str, ...] = ()


@dataclass(frozen=True)
class TemplateInspection:
    entity: str
    path: Path | None
    status: Literal["PASS", "REVIEW"]
    bank_sheet: str | None
    warnings: tuple[str, ...] = ()


def default_competence(reference_date: date) -> str:
    """Return M-1, independent of the execution month."""
    if reference_date.month == 1:
        return f"{reference_date.year - 1}-12"
    return f"{reference_date.year}-{reference_date.month - 1:02d}"


def month_folder(root: str | Path, entity: str, competence: str) -> Path:
    _validate(entity, competence)
    year, month = competence.split("-")
    return Path(root) / year / f"{month}{year}" / entity


def canonical_workbook_path(root: str | Path, entity: str, competence: str) -> Path:
    return month_folder(root, entity, competence) / f"{_NAMES[entity]}_{competence.replace('-', '')}.xlsx"


def find_existing_workbook(root: str | Path, entity: str, competence: str) -> Path | None:
    """Prefer the new canonical name, while recognizing the MDB legacy name."""
    canonical = canonical_workbook_path(root, entity, competence)
    if canonical.is_file():
        return canonical
    if entity == "MDB":
        year, month = competence.split("-")
        legacy = canonical.with_name(f"{_NAMES[entity]}_{month}{year}.xlsx")
        if legacy.is_file():
            return legacy
    return None


def template_path(entity: str) -> Path:
    _validate(entity, "2026-01")
    return TEMPLATE_ROOT / f"{_NAMES[entity]}_TEMPLATE_V1.xlsx"


def inspect_template(entity: str, path: str | Path | None = None) -> TemplateInspection:
    """Accept only an empty operational bank sheet plus persistent structure."""
    candidate = Path(path) if path is not None else template_path(entity)
    if not candidate.is_file():
        return TemplateInspection(entity, None, "REVIEW", None, ("Template limpo ainda não homologado.",))
    try:
        book = load_workbook(candidate, read_only=True, data_only=False, keep_links=True)
        bank_sheet = next((name for name in book.sheetnames if any(token in name.casefold() for token in {"banco", "btg", "safra", "extrato"})), None)
        if bank_sheet is None:
            return TemplateInspection(entity, candidate, "REVIEW", None, ("A estrutura bancária do template não foi identificada.",))
        sheet = book[bank_sheet]
        headers = [str(cell.value or "").strip().casefold() for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
        required = ("data", "descrição", "crédito", "fonte pagadora")
        if not all(value in headers for value in required):
            return TemplateInspection(entity, candidate, "REVIEW", bank_sheet, ("O template não preserva as colunas operacionais exigidas.",))
        if any(any(cell.value is not None for cell in row) for row in sheet.iter_rows(min_row=2)):
            return TemplateInspection(entity, candidate, "REVIEW", bank_sheet, ("O template contém transações bancárias mensais.",))
        return TemplateInspection(entity, candidate, "PASS", bank_sheet)
    except Exception as exc:
        return TemplateInspection(entity, candidate, "REVIEW", None, (f"Não foi possível validar o template: {exc}",))
    finally:
        try:
            book.close()
        except UnboundLocalError:
            pass


def validate_with_native_excel(path: Path) -> tuple[bool, bool, tuple[str, ...]]:
    """Recalculate and reopen a working copy in an isolated Excel process."""
    if sys.platform != "win32":
        return False, True, ("WINDOWS_EXCEL_REQUIRED: validação nativa exige Windows e Excel desktop.",)
    try:
        import win32com.client
        excel = win32com.client.DispatchEx("Excel.Application")
        excel.Visible = False
        excel.DisplayAlerts = False
        try:
            workbook = excel.Workbooks.Open(str(path), 0, False)
            excel.Calculation = -4105  # automatic
            excel.CalculateFullRebuild()
            workbook.Save()
            workbook.Close(True)
            reopened = excel.Workbooks.Open(str(path), 0, True)
            reopened.Close(False)
        finally:
            excel.Quit()
        return True, False, ()
    except Exception as exc:
        return False, True, (str(exc),)


class MonthPreparationService:
    """Create-only preparation service; no UI or mapping rules live here."""

    def __init__(self, *, monthly_root: str | Path | None = None, templates_root: str | Path | None = None,
                 native_excel_validator: Callable[[Path], tuple[bool, bool, tuple[str, ...]]] | None = None) -> None:
        self.monthly_root = Path(monthly_root) if monthly_root is not None else ROOT
        self.templates_root = Path(templates_root) if templates_root is not None else TEMPLATE_ROOT
        self.native_excel_validator = native_excel_validator

    def template_inspection(self, entity: str) -> TemplateInspection:
        return inspect_template(entity, self.templates_root / f"{_NAMES[entity]}_TEMPLATE_V1.xlsx")

    def discover(self, entity: str, competence: str, result: BankExtractionResult | None = None) -> MonthPreparationResult:
        _validate(entity, competence)
        existing = find_existing_workbook(self.monthly_root, entity, competence)
        rows, total = _bank_summary(result)
        folder = month_folder(self.monthly_root, entity, competence)
        if existing:
            return MonthPreparationResult(entity, competence, ENTITY_BANK[entity], "PREPARED", folder, existing, True, False, rows, total)
        template = self.template_inspection(entity)
        return MonthPreparationResult(entity, competence, ENTITY_BANK[entity], "REVIEW" if template.status == "PASS" else "BLOCKED", folder, None, False, False, rows, total, template.warnings)

    def prepare_month(self, entity: str, competence: str, validated_bank_result: BankExtractionResult) -> MonthPreparationResult:
        _validate_result(entity, competence, validated_bank_result)
        before = self.discover(entity, competence, validated_bank_result)
        if before.already_existed:
            return before
        template = self.template_inspection(entity)
        if template.status != "PASS" or template.path is None or template.bank_sheet is None:
            return before
        validator = self.native_excel_validator or validate_with_native_excel
        try:
            with TemporaryDirectory(prefix="muv-month-preparation-") as temporary:
                staged = Path(temporary) / canonical_workbook_path(self.monthly_root, entity, competence).name
                shutil.copy2(template.path, staged)
                self._write_credits(staged, template.bank_sheet, validated_bank_result)
                reopen_ok, repair_required, validation_errors = validator(staged)
                if not reopen_ok or repair_required or validation_errors:
                    return MonthPreparationResult(entity, competence, ENTITY_BANK[entity], "BLOCKED", before.month_folder, None, False, False, before.bank_rows, before.bank_total, errors=validation_errors or ("VALIDAÇÃO_EXCEL_NATIVO_FALHOU",))
                destination = canonical_workbook_path(self.monthly_root, entity, competence)
                publication = _create_only(staged, destination)
                if publication == "CONFLICT":
                    return MonthPreparationResult(entity, competence, ENTITY_BANK[entity], "CONFLICT", before.month_folder, destination if destination.is_file() else None, destination.is_file(), False, before.bank_rows, before.bank_total, errors=("CONFLITO_DE_CRIAÇÃO",))
                if publication != "PASS":
                    return MonthPreparationResult(entity, competence, ENTITY_BANK[entity], "BLOCKED", before.month_folder, None, False, False, before.bank_rows, before.bank_total, errors=("SALVAMENTO_SEGURO_INDISPONÍVEL",))
            return MonthPreparationResult(entity, competence, ENTITY_BANK[entity], "PREPARED", before.month_folder, destination, False, True, before.bank_rows, before.bank_total)
        except Exception as exc:
            return MonthPreparationResult(entity, competence, ENTITY_BANK[entity], "BLOCKED", before.month_folder, None, False, False, before.bank_rows, before.bank_total, errors=(str(exc),))

    @staticmethod
    def _write_credits(path: Path, bank_sheet: str, result: BankExtractionResult) -> None:
        book = load_workbook(path, data_only=False, keep_links=True)
        try:
            sheet = book[bank_sheet]
            for row_number, credit in enumerate(result.operational_credits, start=2):
                for column, value in enumerate((credit.transaction_date, credit.description, credit.credit, credit.payor_source), start=1):
                    sheet.cell(row=row_number, column=column, value=value)
            for cell in sheet["A"][1:]: cell.number_format = "DD/MM/YYYY"
            for cell in sheet["C"][1:]: cell.number_format = "#,##0.00"
            book.save(path)
        finally:
            book.close()


def _create_only(source: Path, destination: Path) -> str:
    """Create a fresh destination atomically; never overwrite an existing month."""
    return create_workbook_only(source, destination).status


def _bank_summary(result: BankExtractionResult | None) -> tuple[int, Decimal]:
    if result is None:
        return 0, Decimal("0.00")
    return len(result.operational_credits), sum((item.credit for item in result.operational_credits), Decimal("0.00"))


def _validate(entity: str, competence: str) -> None:
    if entity not in ENTITY_BANK:
        raise ValueError("Entidade inválida.")
    if len(competence) != 7 or competence[4] != "-" or not competence[:4].isdigit() or not competence[5:].isdigit() or not 1 <= int(competence[5:]) <= 12:
        raise ValueError("Competência inválida.")


def _validate_result(entity: str, competence: str, result: BankExtractionResult) -> None:
    _validate(entity, competence)
    if result.entity != entity or result.period != competence or result.bank_source != ENTITY_BANK[entity]:
        raise ValueError("Resultado bancário não corresponde ao contexto selecionado.")
    if result.errors:
        raise ValueError("Resultado bancário bloqueado.")
