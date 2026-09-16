from decimal import Decimal
from hashlib import sha256
import json
from pathlib import Path
import sys

import pdfplumber
from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from bank_extraction_core import BankExtractionService
from mdb_safra_worksheet_writer import MdbSafraWorksheetWriter, SyntheticOpenpyxlWorkbookBackend


class _Page:
    def __init__(self, period: str = "09/2026") -> None:
        self.period = period

    def extract_text(self, **_):
        return f"""Extrato de Movimentação TESTE
CNPJ: 000
AG: 000 | CONTA: 000-000
Período de 01/{self.period} a 30/{self.period}
LANÇAMENTOS REALIZADOS
01/09 TED RECEBIDA BCO 000 PAGADOR TESTE 0001 1.234,56
complemento sintético"""


class _Pdf:
    def __init__(self, period: str = "09/2026") -> None:
        self.pages = [_Page(period)]

    def __enter__(self): return self
    def __exit__(self, *_): return False


def _template(path: Path, sheet_name: str = "Safra") -> None:
    book = Workbook()
    sheet = book.active
    sheet.title = sheet_name
    sheet.append(["data", "lancamento", "complemento", "documento", "valor_str", "valor", "Fonte Pagadora", "Situação"])
    reconciliation = book.create_sheet("2. Conciliação TESTE")
    reconciliation["A1"] = "Fonte"; reconciliation["B1"] = "Total"
    reconciliation["A2"] = "Fonte Teste"
    reconciliation["B2"] = f'=SUMIFS(\'{sheet_name}\'!$F:$F,\'{sheet_name}\'!$G:$G,A2)'
    staging = book.create_sheet("bs")
    staging.append(["data", "lancamento", "complemento", "documento", "valor_str", "valor"])
    book.create_sheet("Sentinela")["A1"] = "sintetico"
    book.save(path)


def _map(path: Path, aliases: list[dict[str, str]]) -> None:
    path.write_text(json.dumps({"aliases": aliases}), encoding="utf-8")


def test_mdb_prepare_month_uses_real_facade_and_is_create_only(monkeypatch, tmp_path):
    template = tmp_path / "template.xlsx"
    sheet_name = "1. SAFRA TESTE"
    _template(template, sheet_name)
    _map(tmp_path / "map.json", [{"bank_payor_alias": "PAGADOR TESTE", "canonical_royalty_source": "Fonte Teste"}])
    monkeypatch.setenv("MUV_SAFRA_SOURCE_MAP", str(tmp_path / "map.json"))
    monkeypatch.setattr(pdfplumber, "open", lambda _: _Pdf())
    service = BankExtractionService(mdb_template_path=template, monthly_root=tmp_path, mdb_safra_sheet_name=sheet_name, mdb_writer=MdbSafraWorksheetWriter(backend=SyntheticOpenpyxlWorkbookBackend()))

    first_status, first_detail = service.prepare_month(
        entity="MDB", period="2026-09", source_bytes=b"synthetic", source_name="statement.pdf"
    )
    final = tmp_path / "2026" / "092026" / "MDB" / "Conciliação - Músicas do Brasil_202609.xlsx"
    assert first_status == "PREPARED" and Path(first_detail) == final and final.is_file()
    output = load_workbook(final)
    row = output[sheet_name][2]
    assert row[0].value.strftime("%Y-%m-%d") == "2026-09-01"
    assert row[1].value == "TED RECEBIDA BCO 000 PAGADOR TESTE"
    assert all(cell.value is not None for cell in row[2:5])
    assert row[4].value == "1.234,56"
    assert Decimal(str(row[5].value)) == Decimal("1234.56") and row[6].value == "Fonte Teste"
    assert row[7].value is None
    downstream_formula = output["2. Conciliação TESTE"]["B2"].value
    assert sheet_name in downstream_formula and "$F:$F" in downstream_formula and "$G:$G" in downstream_formula
    assert output["bs"].max_row == 1
    assert output["Sentinela"]["A1"].value == "sintetico"
    output.close()
    first_hash = sha256(final.read_bytes()).hexdigest()

    second_status, second_detail = service.prepare_month(
        entity="MDB", period="2026-09", source_bytes=b"synthetic", source_name="statement.pdf"
    )
    assert second_status == "ALREADY_EXISTS" and Path(second_detail) == final
    assert sha256(final.read_bytes()).hexdigest() == first_hash


def test_mdb_prepare_month_blocks_missing_template_without_fallback(tmp_path):
    service = BankExtractionService(mdb_template_path=tmp_path / "missing-template.xlsx", monthly_root=tmp_path)
    status, _ = service.prepare_month(entity="MDB", period="2026-09", source_bytes=b"synthetic", source_name="statement.pdf")
    final = tmp_path / "2026" / "092026" / "MDB" / "Conciliação - Músicas do Brasil_202609.xlsx"
    assert status == "BLOCKED" and not final.exists()


def test_mdb_prepare_month_blocks_review_result(monkeypatch, tmp_path):
    template = tmp_path / "template.xlsx"
    _template(template)
    _map(tmp_path / "map.json", [])
    monkeypatch.setenv("MUV_SAFRA_SOURCE_MAP", str(tmp_path / "map.json"))
    monkeypatch.setattr(pdfplumber, "open", lambda _: _Pdf())
    service = BankExtractionService(mdb_template_path=template, monthly_root=tmp_path)
    status, _ = service.prepare_month(entity="MDB", period="2026-09", source_bytes=b"synthetic", source_name="statement.pdf")
    final = tmp_path / "2026" / "092026" / "MDB" / "Conciliação - Músicas do Brasil_202609.xlsx"
    assert status == "BLOCKED" and not final.exists()


def test_mdb_prepare_month_blocks_competence_mismatch(monkeypatch, tmp_path):
    template = tmp_path / "template.xlsx"
    _template(template)
    _map(tmp_path / "map.json", [{"bank_payor_alias": "PAGADOR TESTE", "canonical_royalty_source": "Fonte Teste"}])
    monkeypatch.setenv("MUV_SAFRA_SOURCE_MAP", str(tmp_path / "map.json"))
    monkeypatch.setattr(pdfplumber, "open", lambda _: _Pdf("08/2026"))
    service = BankExtractionService(mdb_template_path=template, monthly_root=tmp_path)
    status, _ = service.prepare_month(entity="MDB", period="2026-09", source_bytes=b"synthetic", source_name="statement.pdf")
    final = tmp_path / "2026" / "092026" / "MDB" / "Conciliação - Músicas do Brasil_202609.xlsx"
    assert status == "BLOCKED" and not final.exists()


def test_mdb_prepare_month_blocks_missing_configured_safra_sheet(monkeypatch, tmp_path):
    template = tmp_path / "template.xlsx"
    _template(template)
    _map(tmp_path / "map.json", [{"bank_payor_alias": "PAGADOR TESTE", "canonical_royalty_source": "Fonte Teste"}])
    monkeypatch.setenv("MUV_SAFRA_SOURCE_MAP", str(tmp_path / "map.json"))
    monkeypatch.setattr(pdfplumber, "open", lambda _: _Pdf())
    service = BankExtractionService(mdb_template_path=template, monthly_root=tmp_path, mdb_safra_sheet_name="SAFRA AUSENTE")
    status, _ = service.prepare_month(entity="MDB", period="2026-09", source_bytes=b"synthetic", source_name="statement.pdf")
    final = tmp_path / "2026" / "092026" / "MDB" / "Conciliação - Músicas do Brasil_202609.xlsx"
    assert status == "BLOCKED" and not final.exists()


def test_mdb_production_service_blocks_non_windows_without_publication(monkeypatch, tmp_path):
    monkeypatch.setattr(sys, "platform", "darwin")
    template = tmp_path / "template.xlsx"
    _template(template)
    original = template.read_bytes()
    _map(tmp_path / "map.json", [{"bank_payor_alias": "PAGADOR TESTE", "canonical_royalty_source": "Fonte Teste"}])
    monkeypatch.setenv("MUV_SAFRA_SOURCE_MAP", str(tmp_path / "map.json"))
    monkeypatch.setattr(pdfplumber, "open", lambda _: _Pdf())
    service = BankExtractionService(mdb_template_path=template, monthly_root=tmp_path / "blocked-production-months")
    status, detail = service.prepare_month(entity="MDB", period="2026-09", source_bytes=b"synthetic", source_name="statement.pdf")
    assert status == "BLOCKED" and "requires Windows" in detail
    assert not (tmp_path / "blocked-production-months").exists()
    assert template.read_bytes() == original
