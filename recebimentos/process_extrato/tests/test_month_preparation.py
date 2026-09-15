from __future__ import annotations

from datetime import date
from decimal import Decimal
from pathlib import Path
from tempfile import TemporaryDirectory
import sys
import unittest

from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from bank_extraction_core import BankExtractionResult, OperationalCredit
from month_preparation import (MonthPreparationService, canonical_workbook_path,
                               _create_only, default_competence, find_existing_workbook, month_folder)


def fixture_result(entity: str, period: str, *, unknown: bool = False) -> BankExtractionResult:
    source = "REVIEW / UNKNOWN" if unknown else "Fonte aprovada"
    status = "REVIEW" if unknown else "PASS"
    credit = OperationalCredit(date(2026, 9, 3), "Crédito de teste", Decimal("104.52"), source, status)
    return BankExtractionResult(entity, "BTG" if entity == "HM" else "SAFRA", period, "fixture.pdf", "PASS", "PASS", "PASS", 1, 1, 0, Decimal("0"), Decimal("104.52"), Decimal("0"), Decimal("104.52"), (credit,), int(unknown), (), (), "fixture")


def create_clean_template(folder: Path, entity: str) -> None:
    folder.mkdir(parents=True, exist_ok=True)
    name = "Conciliação - Hurst Music" if entity == "HM" else "Conciliação - Músicas do Brasil"
    book = Workbook(); bank = book.active; bank.title = "Banco"
    bank.append(["Data", "Descrição", "Crédito", "Fonte Pagadora"])
    mapping = book.create_sheet("De_Para"); mapping.append(["Descrição", "Fonte Pagadora"]); mapping.append(["Alias", "Fonte"])
    book.save(folder / f"{name}_TEMPLATE_V1.xlsx")


class MonthPreparationTests(unittest.TestCase):
    def test_m1_default_and_year_rollover(self):
        self.assertEqual(default_competence(date(2026, 9, 14)), "2026-08")
        self.assertEqual(default_competence(date(2026, 10, 14)), "2026-09")
        self.assertEqual(default_competence(date(2027, 1, 5)), "2026-12")

    def test_hm_and_mdb_canonical_paths(self):
        root = Path("C:/temp-root")
        self.assertEqual(canonical_workbook_path(root, "HM", "2026-09").name, "Conciliação - Hurst Music_202609.xlsx")
        self.assertEqual(canonical_workbook_path(root, "MDB", "2026-09").name, "Conciliação - Músicas do Brasil_202609.xlsx")
        self.assertEqual(month_folder(root, "MDB", "2026-09"), root / "2026" / "092026" / "MDB")

    def test_existing_month_is_idempotent_and_legacy_mdb_is_detected(self):
        with TemporaryDirectory() as temporary:
            root = Path(temporary)
            legacy = month_folder(root, "MDB", "2026-08") / "Conciliação - Músicas do Brasil_082026.xlsx"
            legacy.parent.mkdir(parents=True); legacy.write_bytes(b"existing")
            found = find_existing_workbook(root, "MDB", "2026-08")
            self.assertEqual(found, legacy)
            service = MonthPreparationService(monthly_root=root, templates_root=root / "templates")
            outcome = service.prepare_month("MDB", "2026-08", fixture_result("MDB", "2026-08"))
            self.assertTrue(outcome.already_existed)
            self.assertFalse(outcome.created_now)

    def test_hm_and_mdb_create_only_in_temporary_root(self):
        with TemporaryDirectory() as temporary:
            root = Path(temporary); templates = root / "templates"
            create_clean_template(templates, "HM"); create_clean_template(templates, "MDB")
            validator = lambda path: (True, False, ())
            service = MonthPreparationService(monthly_root=root / "months", templates_root=templates, native_excel_validator=validator)
            for entity in ("HM", "MDB"):
                outcome = service.prepare_month(entity, "2026-09", fixture_result(entity, "2026-09"))
                self.assertEqual(outcome.status, "PREPARED")
                self.assertTrue(outcome.created_now)
                self.assertTrue(outcome.workbook_path.is_file())
                again = service.prepare_month(entity, "2026-09", fixture_result(entity, "2026-09"))
                self.assertTrue(again.already_existed)
                self.assertFalse(again.created_now)

    def test_inserted_credit_and_unknown_mapping_are_preserved(self):
        with TemporaryDirectory() as temporary:
            root = Path(temporary); templates = root / "templates"; create_clean_template(templates, "HM")
            service = MonthPreparationService(monthly_root=root / "months", templates_root=templates, native_excel_validator=lambda path: (True, False, ()))
            outcome = service.prepare_month("HM", "2026-09", fixture_result("HM", "2026-09", unknown=True))
            book = load_workbook(outcome.workbook_path, read_only=True, data_only=False)
            values = list(book["Banco"].values)
            book.close()
            self.assertEqual(values[1][1:], ("Crédito de teste", 104.52, "REVIEW / UNKNOWN"))

    def test_excel_validation_contract_blocks_creation(self):
        with TemporaryDirectory() as temporary:
            root = Path(temporary); templates = root / "templates"; create_clean_template(templates, "HM")
            service = MonthPreparationService(monthly_root=root / "months", templates_root=templates, native_excel_validator=lambda path: (False, True, ("repair",)))
            outcome = service.prepare_month("HM", "2026-09", fixture_result("HM", "2026-09"))
            self.assertEqual(outcome.status, "BLOCKED")
            self.assertFalse(canonical_workbook_path(root / "months", "HM", "2026-09").exists())

    def test_no_template_blocks_official_creation(self):
        with TemporaryDirectory() as temporary:
            service = MonthPreparationService(monthly_root=Path(temporary) / "months", templates_root=Path(temporary) / "templates")
            outcome = service.prepare_month("HM", "2026-09", fixture_result("HM", "2026-09"))
            self.assertEqual(outcome.status, "BLOCKED")
            self.assertFalse(outcome.created_now)

    def test_create_only_conflict_never_overwrites_destination(self):
        with TemporaryDirectory() as temporary:
            root = Path(temporary); source = root / "source.xlsx"; destination = root / "destination.xlsx"
            source.write_bytes(b"new"); destination.write_bytes(b"existing")
            self.assertEqual(_create_only(source, destination), "CONFLICT")
            self.assertEqual(destination.read_bytes(), b"existing")


if __name__ == "__main__":
    unittest.main()
