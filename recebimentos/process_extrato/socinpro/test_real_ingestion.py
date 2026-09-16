from datetime import date
from decimal import Decimal
import json
from pathlib import Path

import pytest
from openpyxl import Workbook

from socinpro.real_ingestion import (MAPPING_ENVIRONMENT_VARIABLE, SocinproIngestionError, load_mapping_from_environment, load_mapping_file, mapping_resolver, parse_payment_pdf, parse_payment_workbook)
from socinpro.socinpro_vertical import BankReference, process_socinpro


def write_text_pdf(path: Path, lines: list[str]) -> None:
    escaped = [line.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)") for line in lines]
    stream = "BT /F1 12 Tf 72 720 Td " + " ".join(f"({line}) Tj 0 -18 Td" for line in escaped) + " ET"
    objects = [
        "<< /Type /Catalog /Pages 2 0 R >>",
        "<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>",
        f"<< /Length {len(stream.encode('latin-1'))} >>\nstream\n{stream}\nendstream",
        "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    content = b"%PDF-1.4\n"; offsets = [0]
    for number, item in enumerate(objects, start=1):
        offsets.append(len(content)); content += f"{number} 0 obj\n{item}\nendobj\n".encode("latin-1")
    start = len(content); content += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    content += b"".join(f"{offset:010d} 00000 n \n".encode() for offset in offsets[1:])
    content += f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{start}\n%%EOF\n".encode()
    path.write_bytes(content)


def payment_pdf(path: Path) -> None:
    write_text_pdf(path, ["Demonstrativo do Titular", "TITULAR SINTETICO", "Cod. SOCINPRO:", "Cod. ECAD:", "123456", "25/09/2026 R$ -100,50 Pagamento efetuado"])


def mapping_file(path: Path, entity: str = "HM") -> Path:
    path.write_text(json.dumps({"schema_version": 1, "mappings": [{"entity": entity, "source_code": "123456", "titular": "TITULAR SINTETICO", "catalog": "CATALOGO SINTETICO", "deal": "DEAL-SYNTHETIC-001", "active": True}]}), encoding="utf-8")
    return path


def test_real_pdf_layout_is_parsed_with_decimal_and_required_fields(tmp_path):
    path = tmp_path / "socinpro.pdf"; payment_pdf(path)
    parsed = parse_payment_pdf(path)
    assert parsed.source_code == "123456"
    assert parsed.titular == "TITULAR SINTETICO"
    assert parsed.payment_date == date(2026, 9, 25)
    assert parsed.original_value == Decimal("-100.50")
    assert parsed.receipt_value == Decimal("100.50")


def test_pdf_account_line_supersedes_document_subtype_for_titular(tmp_path):
    path = tmp_path / "portal-layout.pdf"
    write_text_pdf(path, ["Demonstrativo do Titular", "Extrato analitico", "HURST MUSIC SPE I LTDA Cod. SOCINPRO: 9396822", "Pagamento efetuado: 24/07/2026 - R$-5.054,23"])
    parsed = parse_payment_pdf(path)
    assert parsed.titular == "HURST MUSIC SPE I LTDA"
    assert parsed.source_code == "9396822"


def test_real_workbook_projection_is_read_only_and_requires_legacy_headers(tmp_path):
    path = tmp_path / "payments.xlsx"; book = Workbook(); sheet = book.active; sheet.title = "pagamentos"
    sheet.append(["cod_socinpro", "titular", "data_pagamento", "valor_pagamento", "arquivo_analitico"])
    sheet.append(["123456", "TITULAR SINTETICO", date(2026, 9, 25), "-100,50", "evidence.pdf"])
    book.save(path); book.close()
    parsed = parse_payment_workbook(path)
    assert len(parsed) == 1 and parsed[0].receipt_value == Decimal("100.50")
    missing = tmp_path / "missing.xlsx"; book = Workbook(); book.active.title = "pagamentos"; book.active.append(["titular"]); book.save(missing); book.close()
    with pytest.raises(SocinproIngestionError, match="CAMPOS_WORKBOOK_AUSENTES"):
        parse_payment_workbook(missing)


def test_operational_payment_workbook_layout_with_header_on_row_four_is_supported(tmp_path):
    path = tmp_path / "operational.xlsx"; book = Workbook(); sheet = book.active; sheet.title = "Pagamentos"
    sheet.append([]); sheet.append(["Pagamentos"]); sheet.append([])
    sheet.append(["Data do pagamento", "Código na associação", "Titular", "Valor", "Documento", "Status", "Observação"])
    sheet.append([date(2026, 9, 25), "123456", "TITULAR SINTETICO", Decimal("-100.50"), "evidence.pdf", "ok", None])
    book.save(path); book.close()
    parsed = parse_payment_workbook(path)
    assert len(parsed) == 1 and parsed[0].source_code == "123456" and parsed[0].receipt_value == Decimal("100.50")


@pytest.mark.parametrize("entity,bank", (("HM", "BTG"), ("MDB", "SAFRA")))
def test_real_parser_flows_to_each_vertical_with_external_mapping(tmp_path, entity, bank):
    source = tmp_path / "socinpro.pdf"; payment_pdf(source)
    mapping = mapping_file(tmp_path / "mapping.json", entity)
    parsed = parse_payment_pdf(source).to_payment(entity=entity, competence="2026-09")
    mappings = load_mapping_from_environment({MAPPING_ENVIRONMENT_VARIABLE: str(mapping)})
    result = process_socinpro(entity=entity, competence="2026-09", payments=(parsed,), bank_result=BankReference(entity, "2026-09", bank, Decimal("100.50"), "synthetic-approved-bank"), catalog_resolver=mapping_resolver(mappings, entity))
    assert result.status == "PASS"


def test_missing_or_ambiguous_mapping_is_blocked_without_fallback(tmp_path):
    with pytest.raises(SocinproIngestionError, match="MAPPING_SOCINPRO_NAO_CONFIGURADO"):
        load_mapping_from_environment({})
    path = tmp_path / "ambiguous.json"
    payload = {"schema_version": 1, "mappings": [{"entity": "HM", "source_code": "123456", "titular": "A", "catalog": "C", "deal": "D", "active": True}, {"entity": "HM", "source_code": "123456", "titular": "B", "catalog": "C", "deal": "D", "active": True}]}
    path.write_text(json.dumps(payload), encoding="utf-8")
    with pytest.raises(SocinproIngestionError, match="MAPPING_SOCINPRO_AMBIGUO"):
        load_mapping_file(path)


def test_invalid_pdf_and_required_fields_fail_closed(tmp_path):
    invalid = tmp_path / "invalid.pdf"; invalid.write_text("not a pdf", encoding="utf-8")
    with pytest.raises(SocinproIngestionError, match="PDF_SOCINPRO_INVALIDO"):
        parse_payment_pdf(invalid)
    positive = tmp_path / "positive.pdf"
    write_text_pdf(positive, ["Demonstrativo do Titular", "TITULAR SINTETICO", "Cod. SOCINPRO: 123456", "25/09/2026 R$ 100,50 Pagamento efetuado"])
    with pytest.raises(SocinproIngestionError, match="SINAL_SOCINPRO_INESPERADO"):
        parse_payment_pdf(positive)
