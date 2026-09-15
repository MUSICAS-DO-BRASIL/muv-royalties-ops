"""Fail-closed SOCINPRO source vertical shared by HM and MDB.

The module deliberately accepts already-extracted source rows. PDF/portal
acquisition stays outside this clean repository until a reviewed, portable
adapter is supplied. This keeps the source contract testable without placing
credentials, mappings, or financial documents in Git.
"""
from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime, timezone
from decimal import Decimal, ROUND_HALF_UP, InvalidOperation
from hashlib import sha256
from pathlib import Path
from tempfile import TemporaryDirectory
from typing import Callable, Iterable, Literal, Mapping
import json
import os
import re

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

from cloud_aware_workbook_publisher import create_workbook_only, semantic_fingerprint


SOURCE_NAME = "SOCINPRO"
ENTITY_BANK = {"HM": "BTG", "MDB": "SAFRA"}
CENT = Decimal("0.01")
Status = Literal["PASS", "REVIEW", "BLOCKED"]


class SocinproContractError(ValueError):
    """Input cannot safely enter the SOCINPRO vertical."""


@dataclass(frozen=True)
class SocinproPayment:
    """Lossless canonical payment contract after source extraction.

    ``original_reference`` is a stable external reference or evidence URI, not
    a local operational path. Monetary values are Decimal by construction.
    """
    entity: str
    competence: str
    payment_date: date
    titular: str
    source_code: str
    gross_value: Decimal
    source_identity: str
    original_reference: str
    document: str = ""
    observation: str = ""

    @property
    def payment_id(self) -> str:
        payload = "|".join((self.entity, self.competence, self.payment_date.isoformat(), self.source_code, self.source_identity, format(self.gross_value, "f")))
        return sha256(payload.encode("utf-8")).hexdigest()


@dataclass(frozen=True)
class CatalogRelation:
    catalog: str
    deal: str
    status: Literal["DETERMINISTIC", "REVIEW"] = "DETERMINISTIC"


@dataclass(frozen=True)
class BankReference:
    entity: str
    competence: str
    bank_source: str
    total: Decimal
    reference: str = ""


@dataclass(frozen=True)
class SocinproLine:
    payment: SocinproPayment
    catalog: CatalogRelation | None
    status: Status
    warning: str = ""


@dataclass(frozen=True)
class SocinproResult:
    entity: str
    competence: str
    bank_source: str
    status: Status
    documentary_total: Decimal
    bank_total: Decimal
    difference: Decimal
    lines: tuple[SocinproLine, ...]
    pending: tuple[str, ...]
    warnings: tuple[str, ...]
    evidence: tuple[str, ...]
    source_folder: Path | None = None
    demonstrative_path: Path | None = None


def parse_payment_row(row: Mapping[str, object]) -> SocinproPayment:
    """Normalize a source row; no field is inferred from another field."""
    try:
        entity = str(row["entity"]).upper().strip()
        competence = str(row["competence"]).strip()
        payment_date = _date(row["payment_date"])
        gross_value = _decimal(row["gross_value"])
        payment = SocinproPayment(entity, competence, payment_date, str(row["titular"]).strip(), str(row["source_code"]).strip(), gross_value, str(row["source_identity"]).strip(), str(row["original_reference"]).strip(), str(row.get("document") or "").strip(), str(row.get("observation") or "").strip())
    except KeyError as exc:
        raise SocinproContractError(f"CAMPO_SOCINPRO_AUSENTE:{exc.args[0]}") from exc
    _validate_payment(payment)
    return payment


def process_socinpro(
    *, entity: str, competence: str, payments: Iterable[SocinproPayment], bank_result: BankReference,
    catalog_resolver: Callable[[SocinproPayment], CatalogRelation | None],
) -> SocinproResult:
    """Reconcile one entity/month; uncertain catalog associations stay REVIEW."""
    entity = entity.upper().strip()
    _validate_context(entity, competence)
    if bank_result.entity != entity or bank_result.competence != competence or bank_result.bank_source != ENTITY_BANK[entity]:
        raise SocinproContractError("RESULTADO_BANCARIO_INCOMPATIVEL")
    rows = tuple(payments)
    ids: set[str] = set()
    lines: list[SocinproLine] = []
    pending: list[str] = []
    for payment in rows:
        _validate_payment(payment)
        if payment.entity != entity or payment.competence != competence:
            raise SocinproContractError("PAGAMENTO_FORA_DO_CONTEXTO")
        if payment.payment_id in ids:
            raise SocinproContractError("PAGAMENTO_DUPLICADO")
        ids.add(payment.payment_id)
        relation = catalog_resolver(payment)
        if relation is None or relation.status != "DETERMINISTIC" or not relation.catalog.strip() or not relation.deal.strip():
            message = f"RELACAO_CATALOGO_PENDENTE:{payment.source_identity}"
            lines.append(SocinproLine(payment, relation, "REVIEW", message)); pending.append(message)
        else:
            lines.append(SocinproLine(payment, relation, "PASS"))
    documentary_total = sum((line.payment.gross_value for line in lines), Decimal("0.00"))
    difference = (documentary_total - bank_result.total).quantize(CENT, rounding=ROUND_HALF_UP)
    if difference != Decimal("0.00"):
        pending.append(f"DIFERENCA_BANCO:{difference}")
    status: Status = "PASS" if not pending else "REVIEW"
    return SocinproResult(entity, competence, ENTITY_BANK[entity], status, documentary_total, bank_result.total, difference, tuple(lines), tuple(pending), (), tuple(line.payment.original_reference for line in lines))


def publish_operational_result(result: SocinproResult, operational_root: str | Path, technical_root: str | Path) -> SocinproResult:
    """Create only the useful workbook in ROOT/YYYY/MMYYYY/ENTITY/SOCINPRO.

    A differing existing demonstrative is never overwritten: it may contain a
    human review. Technical provenance is written outside the operational tree.
    """
    _validate_context(result.entity, result.competence)
    if result.status != "PASS":
        raise SocinproContractError("PUBLICACAO_SOCINPRO_BLOQUEADA:RESULTADO_EM_REVISAO")
    source_folder = source_folder_for(operational_root, result.entity, result.competence)
    technical = Path(technical_root).resolve()
    if _overlaps(source_folder.resolve(), technical):
        raise SocinproContractError("TECHNICAL_ROOT_DEVE_SER_EXTERNO_A_PASTA_OPERACIONAL")
    filename = f"Demonstrativo_SOCINPRO_{result.competence.replace('-', '')}.xlsx"
    destination = source_folder / filename
    with TemporaryDirectory(prefix="muv-socinpro-") as temporary:
        staged = Path(temporary) / filename
        _write_demonstrative(staged, result)
        if destination.exists():
            if not _same_demonstrative(destination, staged):
                raise SocinproContractError("DEMONSTRATIVO_EXISTENTE_DIVERGENTE")
        else:
            publication = create_workbook_only(staged, destination)
            if publication.status != "PASS":
                raise SocinproContractError("PUBLICACAO_DEMONSTRATIVO_BLOQUEADA:" + ",".join(publication.errors))
    _write_audit(technical, result, destination)
    return SocinproResult(**{**result.__dict__, "source_folder": source_folder, "demonstrative_path": destination})


def source_folder_for(root: str | Path, entity: str, competence: str) -> Path:
    _validate_context(entity, competence)
    year, month = competence.split("-")
    return Path(root) / year / f"{month}{year}" / entity / SOURCE_NAME


def _write_demonstrative(path: Path, result: SocinproResult) -> None:
    book = Workbook(); summary = book.active; summary.title = "Resumo"
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    summary.append(["Demonstrativo SOCINPRO"])
    for row in (("Entidade", result.entity), ("Competência", result.competence), ("Fonte", SOURCE_NAME), ("Quantidade", len(result.lines)), ("Bruto documental", result.documentary_total), ("Ajustes documentados", Decimal("0.00")), ("Líquido esperado", result.documentary_total), ("Banco", result.bank_source), ("Total bancário", result.bank_total), ("Diferença", result.difference), ("Pendências", len(result.pending)), ("Status geral", result.status)):
        summary.append(row)
    payments = book.create_sheet("Pagamentos")
    payments.append(["Data do pagamento", "Código na fonte/associação", "Titular", "Valor", "Documento", "Status", "Observação"])
    for line in result.lines:
        p = line.payment
        payments.append([p.payment_date, p.source_code, p.titular, p.gross_value, p.document, line.status, line.warning or p.observation])
    if result.pending:
        pending = book.create_sheet("Pendencias")
        pending.append(["Pendência"])
        for item in result.pending: pending.append([item])
    for sheet in book.worksheets:
        for cell in sheet[1]: cell.fill = header_fill; cell.font = header_font
        sheet.freeze_panes = "A2"
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = min(max(len(str(cell.value or "")) for cell in column) + 2, 45)
    for cell in summary["B"]:
        if isinstance(cell.value, Decimal): cell.number_format = '#,##0.00'
    for cell in payments["A"][1:]: cell.number_format = "DD/MM/YYYY"
    for cell in payments["D"][1:]: cell.number_format = '#,##0.00'
    book.save(path); book.close()


def _write_audit(root: Path, result: SocinproResult, demonstrative: Path) -> None:
    root.mkdir(parents=True, exist_ok=True)
    payload = {"pipeline_version": "socinpro-v1", "processed_at": datetime.now(timezone.utc).isoformat(), "entity": result.entity, "competence": result.competence, "status": result.status, "documentary_total": str(result.documentary_total), "bank_total": str(result.bank_total), "difference": str(result.difference), "payment_ids": [line.payment.payment_id for line in result.lines], "evidence": list(result.evidence), "pending": list(result.pending), "demonstrative_sha256": _binary_hash(demonstrative)}
    target = root / f"socinpro_{result.entity.lower()}_{result.competence.replace('-', '')}.json"
    temporary = target.with_suffix(".json.tmp"); temporary.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8"); temporary.replace(target)


def _validate_payment(payment: SocinproPayment) -> None:
    _validate_context(payment.entity, payment.competence)
    if not all((payment.titular, payment.source_code, payment.source_identity, payment.original_reference)):
        raise SocinproContractError("IDENTIDADE_SOCINPRO_INCOMPLETA")
    if payment.gross_value <= Decimal("0"):
        raise SocinproContractError("VALOR_SOCINPRO_INVALIDO")


def _validate_context(entity: str, competence: str) -> None:
    if entity not in ENTITY_BANK: raise SocinproContractError("ENTIDADE_SOCINPRO_INVALIDA")
    if not re.fullmatch(r"\d{4}-(0[1-9]|1[0-2])", competence): raise SocinproContractError("COMPETENCIA_SOCINPRO_INVALIDA")


def _decimal(value: object) -> Decimal:
    if isinstance(value, float): raise SocinproContractError("FLOAT_NAO_PERMITIDO")
    try: return Decimal(str(value))
    except (InvalidOperation, ValueError) as exc: raise SocinproContractError("VALOR_SOCINPRO_INVALIDO") from exc


def _date(value: object) -> date:
    if isinstance(value, date): return value
    try: return date.fromisoformat(str(value))
    except ValueError as exc: raise SocinproContractError("DATA_SOCINPRO_INVALIDA") from exc


def _binary_hash(path: Path) -> str:
    return sha256(path.read_bytes()).hexdigest()


def _same_demonstrative(left: Path, right: Path) -> bool:
    if _binary_hash(left) == _binary_hash(right):
        return True
    try:
        return semantic_fingerprint(left) == semantic_fingerprint(right)
    except Exception:
        return False


def _overlaps(left: Path, right: Path) -> bool:
    try: left.relative_to(right); return True
    except ValueError:
        try: right.relative_to(left); return True
        except ValueError: return False
