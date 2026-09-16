from datetime import date
from decimal import Decimal

from openpyxl import load_workbook
import pytest

from socinpro.demonstrativo import DemonstrativoRow, generate_demonstrativo


def rows(count=2):
    return tuple(DemonstrativoRow(date(2026, 8, 25), str(index), f"Titular {index % 2}", f"Artista {index % 2}", Decimal("100.50"), f"{index}.pdf", "MATCHED") for index in range(count))


@pytest.mark.parametrize("entity", ["HM", "MDB"])
def test_generic_demonstrativo_supports_entities_and_totals(tmp_path, entity):
    output = generate_demonstrativo(entity=entity, source="SOCINPRO", competence="2026-08", financial_rows=rows(), review_items=("review",), provenance=("synthetic",), destination=tmp_path / "report.xlsx")
    book = load_workbook(output, data_only=True)
    assert book.sheetnames == ["Pagamentos", "Resumo por Titular", "Pendências"]
    assert book["Pagamentos"].max_row == 3
    assert sum(row[4] for row in book["Pagamentos"].iter_rows(min_row=2, values_only=True)) == Decimal("201.00")
    assert sum(row[2] for row in book["Resumo por Titular"].iter_rows(min_row=2, values_only=True)) == Decimal("201.00")
    book.close()


def test_86_rows_are_not_doubled(tmp_path):
    output = generate_demonstrativo(entity="HM", source="SOCINPRO", competence="2026-08", financial_rows=rows(86), review_items=(), provenance=(), destination=tmp_path / "report.xlsx")
    book = load_workbook(output, data_only=True)
    assert book["Pagamentos"].max_row - 1 == 86
    assert sum(row[4] for row in book["Pagamentos"].iter_rows(min_row=2, values_only=True)) == Decimal("8643.00")
    book.close()
