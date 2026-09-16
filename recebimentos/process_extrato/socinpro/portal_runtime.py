"""Credential preflight and safe orchestration boundary for SOCINPRO portal runs.

This module intentionally does not know about ``LOGINS.xlsx``.  That workbook
is a global reference used by downstream business processes, not a portal
credential source.  Browser automation is injected by the caller so this
module can enforce the credential contract without storing browser profiles,
cookies, or secrets in the repository.
"""
from __future__ import annotations

from dataclasses import dataclass, field
from hashlib import sha256
import logging
import os
from pathlib import Path
import re
from typing import Callable, Iterable, Mapping, Protocol

from openpyxl import load_workbook


LOGGER = logging.getLogger(__name__)
CREDENTIAL_PATH_ENVIRONMENTS = {
    "HM": "MUV_SOCINPRO_HM_CREDENTIALS_PATH",
    "MDB": "MUV_SOCINPRO_MDB_CREDENTIALS_PATH",
}
USER_HEADERS = {"usuario", "login", "user", "username"}
PASSWORD_HEADERS = {"senha", "password", "pass"}


class SocinproPortalRuntimeError(ValueError):
    """Fail-closed error that never includes a credential value."""


class HumanInterventionRequired(RuntimeError):
    """Raised by a browser adapter when CAPTCHA or MFA needs an operator."""


@dataclass(frozen=True)
class RuntimeAccount:
    index: int
    identifier: str
    password: str = field(repr=False)

    @property
    def masked_identifier(self) -> str:
        return "id#" + sha256(self.identifier.encode("utf-8")).hexdigest()[:10]


@dataclass(frozen=True)
class RuntimeCredentialPreflight:
    entity: str
    source_path: Path
    accounts: tuple[RuntimeAccount, ...]

    @property
    def account_count(self) -> int:
        return len(self.accounts)

    @property
    def unique_account_count(self) -> int:
        return len({account.identifier for account in self.accounts})

    @property
    def duplicate_runtime_accounts(self) -> int:
        return self.account_count - self.unique_account_count

    @property
    def download_allowed(self) -> bool:
        return self.account_count > 0 and self.duplicate_runtime_accounts == 0


class PortalAccountSession(Protocol):
    """A browser adapter. Implementations must keep CAPTCHA/MFA human-led."""

    def authenticate(self, account: RuntimeAccount) -> None: ...

    def select_competence(self, competence: str) -> None: ...

    def download_statements(self, account: RuntimeAccount) -> Iterable[Path]: ...


@dataclass(frozen=True)
class AccountDownloadResult:
    account_index: int
    account_masked_identifier: str
    status: str
    downloaded_files: tuple[Path, ...] = ()


def preflight_runtime_credentials(
    entity: str,
    *,
    environment: Mapping[str, str] | None = None,
    credential_path: str | Path | None = None,
) -> RuntimeCredentialPreflight:
    """Read one entity's external credential source and validate it fail-closed.

    ``credential_path`` exists for an approved launcher or test. Production
    launchers should use the entity-specific environment variable. No default
    filesystem location and no cross-entity fallback are permitted.
    """
    normalized_entity = _normalize_entity(entity)
    source = _resolve_source(normalized_entity, environment, credential_path)
    if not source.is_file():
        raise SocinproPortalRuntimeError("RUNTIME_CREDENTIAL_FILE_NOT_FOUND")

    accounts = _read_accounts(source)
    identifiers = [account.identifier for account in accounts]
    if len(identifiers) != len(set(identifiers)):
        raise SocinproPortalRuntimeError("DUPLICATE_RUNTIME_ACCOUNT")

    preflight = RuntimeCredentialPreflight(normalized_entity, source, tuple(accounts))
    for account in preflight.accounts:
        LOGGER.info(
            "SOCINPRO portal credential preflight entity=%s account_index=%s identifier=%s status=ready",
            normalized_entity,
            account.index,
            account.masked_identifier,
        )
    return preflight


def run_portal_download(
    preflight: RuntimeCredentialPreflight,
    competence: str,
    session_factory: Callable[[], PortalAccountSession],
) -> tuple[AccountDownloadResult, ...]:
    """Execute sequentially through an injected browser adapter after preflight.

    This function never creates a browser itself. A caller that encounters MFA
    or CAPTCHA must surface ``HUMAN_INTERVENTION_REQUIRED`` rather than bypass
    it. It is deliberately not invoked by credential preflight.
    """
    if not preflight.download_allowed:
        raise SocinproPortalRuntimeError("PORTAL_DOWNLOAD_NOT_ALLOWED")
    if not re.fullmatch(r"\d{4}-\d{2}", competence):
        raise SocinproPortalRuntimeError("COMPETENCE_INVALID")
    results: list[AccountDownloadResult] = []
    for account in preflight.accounts:
        try:
            session: PortalAccountSession = session_factory()
            session.authenticate(account)
            session.select_competence(competence)
            files = tuple(session.download_statements(account))
            results.append(AccountDownloadResult(account.index, account.masked_identifier, "DOWNLOADED", files))
        except HumanInterventionRequired:
            results.append(AccountDownloadResult(account.index, account.masked_identifier, "HUMAN_INTERVENTION_REQUIRED"))
            break
    return tuple(results)


def _normalize_entity(entity: str) -> str:
    normalized = str(entity).upper().strip()
    if normalized not in CREDENTIAL_PATH_ENVIRONMENTS:
        raise SocinproPortalRuntimeError("SOCINPRO_ENTITY_INVALID")
    return normalized


def _resolve_source(
    entity: str,
    environment: Mapping[str, str] | None,
    credential_path: str | Path | None,
) -> Path:
    if credential_path is not None:
        return Path(credential_path).expanduser().resolve()
    values = os.environ if environment is None else environment
    configured = str(values.get(CREDENTIAL_PATH_ENVIRONMENTS[entity]) or "").strip()
    if not configured:
        raise SocinproPortalRuntimeError("RUNTIME_CREDENTIAL_PATH_NOT_CONFIGURED")
    return Path(configured).expanduser().resolve()


def _read_accounts(source: Path) -> list[RuntimeAccount]:
    workbook = load_workbook(source, read_only=True, data_only=True)
    try:
        for sheet in workbook.worksheets:
            for header_row, values in enumerate(sheet.iter_rows(values_only=True), start=1):
                headers = [_header(value) for value in values]
                user_index = next((index for index, name in enumerate(headers) if name in USER_HEADERS), None)
                password_index = next((index for index, name in enumerate(headers) if name in PASSWORD_HEADERS), None)
                if user_index is None or password_index is None:
                    continue
                return _accounts_from_sheet(sheet, header_row, user_index, password_index)
    finally:
        workbook.close()
    raise SocinproPortalRuntimeError("RUNTIME_CREDENTIAL_COLUMNS_MISSING")


def _accounts_from_sheet(sheet, header_row: int, user_index: int, password_index: int) -> list[RuntimeAccount]:
    accounts: list[RuntimeAccount] = []
    for row_number, values in enumerate(sheet.iter_rows(min_row=header_row + 1, values_only=True), start=header_row + 1):
        user = _value_at(values, user_index)
        password = _value_at(values, password_index)
        if not user and not password:
            continue
        if not user or not password:
            raise SocinproPortalRuntimeError(f"RUNTIME_CREDENTIAL_FIELD_MISSING:ROW_{row_number}")
        accounts.append(RuntimeAccount(len(accounts) + 1, user.casefold(), password))
    if not accounts:
        raise SocinproPortalRuntimeError("RUNTIME_CREDENTIALS_EMPTY")
    return accounts


def _header(value: object) -> str:
    return str(value or "").strip().casefold()


def _value_at(values: tuple[object, ...], index: int) -> str:
    value = values[index] if index < len(values) else None
    return str(value).strip() if value is not None else ""
