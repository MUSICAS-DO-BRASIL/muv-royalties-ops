from datetime import datetime
from decimal import Decimal
from openpyxl import Workbook
import pytest

from socinpro.official_workbook_adapter import OfficialWorkbookAdapter, WorkbookSchemaError, resolved_source


def workbook(path):
    book=Workbook(); bank=book.active; bank.title="BTG"
    bank.append(["Data","Descrição","Crédito","Fonte Pagadora"])
    bank.append([datetime(2026,8,1),"SOCINPRO",100.5,"SOCINPRO"])
    bank.append([datetime(2026,9,1),"OTHER",1,"OTHER"])
    mapping=book.create_sheet("de_para_fontes"); mapping.append(["de_descricao_extrato","para_fonte_pagadora"]); mapping.append(["SOCINPRO","SOCINPRO"])
    book.save(path); book.close()


def test_read_only_btg_and_source_mapping(tmp_path):
    path=tmp_path/"official.xlsx"; workbook(path); adapter=OfficialWorkbookAdapter(path)
    assert adapter.load_source_depara()=={"socinpro":"SOCINPRO"}
    assert adapter.load_bank_receipts("2026-08")[0].value==Decimal("100.5")
    receipt=adapter.load_bank_receipts("2026-08")[0]
    assert resolved_source(receipt, adapter.load_source_depara()) == "SOCINPRO"


def test_schema_fails_closed(tmp_path):
    path=tmp_path/"bad.xlsx"; Workbook().save(path)
    with pytest.raises(WorkbookSchemaError): OfficialWorkbookAdapter(path).load_bank_receipts("2026-08")
