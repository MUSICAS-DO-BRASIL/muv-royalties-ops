"""Read-only adapter for HM operational workbooks; never writes the source."""
from __future__ import annotations

from dataclasses import dataclass
from datetime import date
from decimal import Decimal
from pathlib import Path
import hashlib
import unicodedata

from openpyxl import load_workbook


class WorkbookSchemaError(ValueError):
    pass


@dataclass(frozen=True)
class BankReceipt:
    payment_date: date
    description: str
    value: Decimal
    source: str


def _key(value: object) -> str:
    value = unicodedata.normalize("NFKD", str(value or ""))
    return "".join(char for char in value if not unicodedata.combining(char)).strip().casefold()


class OfficialWorkbookAdapter:
    """Typed, fail-closed reader for the validated HM workbook schema."""
    def __init__(self, path: str | Path):
        self.path = Path(path)
        if not self.path.is_file():
            raise WorkbookSchemaError("OFFICIAL_WORKBOOK_NOT_FOUND")

    @property
    def sha256(self) -> str:
        return hashlib.sha256(self.path.read_bytes()).hexdigest()

    def _book(self):
        return load_workbook(self.path, read_only=True, data_only=True, keep_links=True)

    @staticmethod
    def _sheet(book, name: str):
        if name not in book.sheetnames:
            raise WorkbookSchemaError(f"WORKSHEET_REQUIRED:{name}")
        return book[name]

    def load_source_depara(self) -> dict[str, str]:
        book = self._book()
        try:
            sheet = self._sheet(book, "de_para_fontes")
            headers = {_key(value): index for index, value in enumerate(next(sheet.iter_rows(values_only=True)))}
            required = ("de_descricao_extrato", "para_fonte_pagadora")
            if any(item not in headers for item in required):
                raise WorkbookSchemaError("SOURCE_DEPARA_COLUMNS_REQUIRED")
            result: dict[str, str] = {}
            for row in sheet.iter_rows(min_row=2, values_only=True):
                left, right = str(row[headers[required[0]]] or "").strip(), str(row[headers[required[1]]] or "").strip()
                if not left or not right: continue
                key = _key(left)
                if key in result and result[key] != right: raise WorkbookSchemaError("SOURCE_DEPARA_AMBIGUOUS")
                result[key] = right
            if not result: raise WorkbookSchemaError("SOURCE_DEPARA_EMPTY")
            return dict(sorted(result.items()))
        finally: book.close()

    def load_bank_receipts(self, competence: str) -> tuple[BankReceipt, ...]:
        book = self._book()
        try:
            name = next((item for item in book.sheetnames if "btg" in _key(item)), None)
            if not name: raise WorkbookSchemaError("BTG_SHEET_REQUIRED")
            sheet = book[name]; headers = {_key(value): index for index, value in enumerate(next(sheet.iter_rows(values_only=True)))}
            required = ("data", "descricao", "credito", "fonte pagadora")
            if any(item not in headers for item in required): raise WorkbookSchemaError("BTG_COLUMNS_REQUIRED")
            year, month = map(int, competence.split("-")); receipts=[]
            for row in sheet.iter_rows(min_row=2, values_only=True):
                raw=row[headers["data"]]
                if not isinstance(raw, date) or raw.year != year or raw.month != month: continue
                value=row[headers["credito"]]
                if value is None: continue
                amount=Decimal(str(value))
                if amount <= 0: continue
                receipts.append(BankReceipt(raw, str(row[headers["descricao"]] or "").strip(), amount, str(row[headers["fonte pagadora"]] or "").strip()))
            return tuple(sorted(receipts, key=lambda item:(item.payment_date,item.description,item.value)))
        finally: book.close()
