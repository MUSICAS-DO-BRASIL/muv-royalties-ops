from datetime import date
from decimal import Decimal
from pathlib import Path
from tempfile import TemporaryDirectory
import json
from hashlib import sha256

import pytest
from openpyxl import load_workbook
from streamlit.testing.v1 import AppTest

from socinpro.socinpro_vertical import BankReference, CatalogRelation, SocinproContractError, SocinproPayment, parse_payment_row, process_socinpro, publish_operational_result, source_folder_for


def payment(entity: str, competence: str, *, identity: str = "SOC-1") -> SocinproPayment:
    return SocinproPayment(entity, competence, date(2026, 9, 3), "Titular sintético", "123", Decimal("100.00"), identity, f"synthetic://{identity}", "DOC-1")


@pytest.mark.parametrize(("entity", "bank"), (("HM", "BTG"), ("MDB", "SAFRA")))
def test_synthetic_vertical_for_hm_and_mdb(entity, bank):
    with TemporaryDirectory() as folder:
        root = Path(folder)
        result = process_socinpro(entity=entity, competence="2026-09", payments=(payment(entity, "2026-09"),), bank_result=BankReference(entity, "2026-09", bank, Decimal("100.00"), "synthetic-bank"), catalog_resolver=lambda _: CatalogRelation("Catálogo sintético", "DEAL-1"))
        published = publish_operational_result(result, root / "operational", root / "technical")
        assert published.status == "PASS"
        assert published.source_folder == source_folder_for(root / "operational", entity, "2026-09")
        assert published.demonstrative_path.is_file()
        assert not any(item.suffix == ".json" for item in published.source_folder.iterdir())
        book = load_workbook(published.demonstrative_path, read_only=True, data_only=False)
        assert book.sheetnames == ["Resumo", "Pagamentos"]
        assert book["Pagamentos"].max_row == 2
        book.close()
        again = publish_operational_result(result, root / "operational", root / "technical")
        assert again.demonstrative_path == published.demonstrative_path


def test_unknown_catalog_relation_is_review_visible_and_blocks_publication():
    with TemporaryDirectory() as folder:
        root = Path(folder)
        result = process_socinpro(entity="HM", competence="2026-09", payments=(payment("HM", "2026-09"),), bank_result=BankReference("HM", "2026-09", "BTG", Decimal("100.00")), catalog_resolver=lambda _: None)
        assert result.status == "REVIEW"
        assert result.pending == ("RELACAO_CATALOGO_PENDENTE:SOC-1",)
        with pytest.raises(SocinproContractError, match="PUBLICACAO_SOCINPRO_BLOQUEADA"):
            publish_operational_result(result, root / "operational", root / "technical")
        assert not source_folder_for(root / "operational", "HM", "2026-09").exists()


def test_fail_closed_for_duplicate_or_incompatible_bank_or_existing_human_file():
    item = payment("MDB", "2026-09")
    with pytest.raises(SocinproContractError, match="PAGAMENTO_DUPLICADO"):
        process_socinpro(entity="MDB", competence="2026-09", payments=(item, item), bank_result=BankReference("MDB", "2026-09", "SAFRA", Decimal("200")), catalog_resolver=lambda _: CatalogRelation("C", "D"))
    with pytest.raises(SocinproContractError, match="RESULTADO_BANCARIO_INCOMPATIVEL"):
        process_socinpro(entity="MDB", competence="2026-09", payments=(item,), bank_result=BankReference("MDB", "2026-09", "BTG", Decimal("100")), catalog_resolver=lambda _: CatalogRelation("C", "D"))
    with TemporaryDirectory() as folder:
        root = Path(folder); result = process_socinpro(entity="MDB", competence="2026-09", payments=(item,), bank_result=BankReference("MDB", "2026-09", "SAFRA", Decimal("100")), catalog_resolver=lambda _: CatalogRelation("C", "D"))
        destination = source_folder_for(root / "operational", "MDB", "2026-09") / "Demonstrativo_SOCINPRO_202609.xlsx"
        destination.parent.mkdir(parents=True); destination.write_bytes(b"human review")
        with pytest.raises(SocinproContractError, match="DEMONSTRATIVO_EXISTENTE_DIVERGENTE"):
            publish_operational_result(result, root / "operational", root / "technical")
        assert destination.read_bytes() == b"human review"


def test_contract_rejects_float_and_missing_identity():
    with pytest.raises(SocinproContractError, match="FLOAT_NAO_PERMITIDO"):
        parse_payment_row({"entity": "HM", "competence": "2026-09", "payment_date": "2026-09-03", "titular": "T", "source_code": "1", "gross_value": 1.0, "source_identity": "I", "original_reference": "evidence://1"})


def test_operator_ui_smoke():
    app = AppTest.from_file(Path(__file__).resolve().parents[1] / "socinpro_operator_app_v1.py", default_timeout=30).run()
    assert not app.exception
    assert app.button[0].disabled


def approved_result():
    return process_socinpro(entity='HM', competence='2026-09',
        payments=(payment('HM', '2026-09'),),
        bank_result=BankReference('HM', '2026-09', 'BTG', Decimal('100.00')),
        catalog_resolver=lambda _: CatalogRelation('Synthetic catalog', 'DEAL-1'))


def test_repeat_preserves_workbook_and_all_previous_audits(tmp_path):
    operational, technical = tmp_path / 'months', tmp_path / 'technical'
    technical.mkdir()
    legacy = technical / 'socinpro_hm_202609.json'
    legacy.write_text('legacy audit', encoding='utf-8')
    result = approved_result()
    first = publish_operational_result(result, operational, technical)
    workbook_bytes = first.demonstrative_path.read_bytes()
    first_audit = next(technical.glob('socinpro_*/audit.json'))
    audit_bytes = first_audit.read_bytes()
    publish_operational_result(result, operational, technical)
    audits = list(technical.glob('socinpro_*/audit.json'))
    assert len(audits) == 2
    assert legacy.read_text() == 'legacy audit'
    assert first_audit.read_bytes() == audit_bytes
    assert first.demonstrative_path.read_bytes() == workbook_bytes
    for path in audits:
        assert json.loads(path.read_text())['demonstrative_sha256'] == sha256(workbook_bytes).hexdigest()
    assert not list(technical.rglob('*.tmp'))
    assert not list(operational.rglob('*.json'))


@pytest.mark.parametrize('location', ['inside', 'same', 'ancestor'])
def test_audit_must_be_outside_the_entire_operational_tree(tmp_path, location):
    operational = tmp_path / 'operational'
    technical = {'inside': operational / 'other-month' / 'audit',
                 'same': operational, 'ancestor': tmp_path}[location]
    with pytest.raises(SocinproContractError, match='TECHNICAL_ROOT_DEVE_SER_EXTERNO'):
        publish_operational_result(approved_result(), operational, technical)
    assert not operational.exists()


def test_audit_failure_reports_preserved_workbook_and_allows_retry(tmp_path):
    operational, technical = tmp_path / 'months', tmp_path / 'technical'
    technical.write_bytes(b'file blocking audit directory')
    with pytest.raises(SocinproContractError, match='AUDITORIA_SOCINPRO_FALHOU:DEMONSTRATIVO_PRESERVADO'):
        publish_operational_result(approved_result(), operational, technical)
    workbook = next(operational.rglob('*.xlsx'))
    original_bytes = workbook.read_bytes()
    technical.unlink()
    publish_operational_result(approved_result(), operational, technical)
    assert workbook.read_bytes() == original_bytes
    assert len(list(technical.glob('socinpro_*/audit.json'))) == 1
