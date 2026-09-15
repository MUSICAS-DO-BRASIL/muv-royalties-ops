"""Portable, entity-agnostic ingestion for reviewed SOCINPRO source files.

This module only reads manually supplied SOCINPRO documents. Portal access,
credentials and monthly publication deliberately remain outside its boundary.
"""
from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from hashlib import sha256
import json
import os
from pathlib import Path
import re
import unicodedata
from typing import Callable, Iterable, Mapping

import pdfplumber
from openpyxl import load_workbook

from .socinpro_vertical import CatalogRelation, SocinproContractError, SocinproPayment


MAPPING_ENVIRONMENT_VARIABLE = "MUV_SOCINPRO_MAPPING_PATH"
SUPPORTED_INPUT_TYPES = ("SOCINPRO_PAYMENT_PDF", "SOCINPRO_PAYMENT_WORKBOOK")


class SocinproIngestionError(ValueError):
    """A source file or external mapping cannot safely be used."""


@dataclass(frozen=True)
class ParsedSocinproPayment:
    source_code: str
    titular: str
    payment_date: date
    original_value: Decimal
    receipt_value: Decimal
    document: str
    source_reference: str

    def to_payment(self, *, entity: str, competence: str) -> SocinproPayment:
        return SocinproPayment(
            entity=entity,
            competence=competence,
            payment_date=self.payment_date,
            titular=self.titular,
            source_code=self.source_code,
            gross_value=self.receipt_value,
            source_identity=_semantic_identity(self.source_code, self.titular, self.payment_date, self.original_value),
            original_reference=self.source_reference,
            document=self.document,
        )


@dataclass(frozen=True)
class SocinproMapping:
    entity: str
    source_code: str
    titular: str
    catalog: str
    deal: str
    active: bool


def parse_payment_pdf(path: str | Path, *, source_reference: str | None = None) -> ParsedSocinproPayment:
    """Parse the text-based SOCINPRO demonstrativo layout recovered from legacy.

    Only the financial ``Pagamento efetuado`` layout is accepted. Detailed
    distribution statements and PDFs without text remain unsupported by design.
    """
    candidate = Path(path)
    if candidate.suffix.casefold() != ".pdf" or not candidate.is_file():
        raise SocinproIngestionError("PDF_SOCINPRO_INVALIDO")
    try:
        with pdfplumber.open(candidate) as document:
            text = "\n".join(page.extract_text() or "" for page in document.pages)
    except Exception as exc:
        raise SocinproIngestionError("PDF_SOCINPRO_INVALIDO") from exc
    if not text.strip():
        raise SocinproIngestionError("PDF_SOCINPRO_SEM_TEXTO")
    return _parse_payment_text(text, source_reference or candidate.as_uri(), candidate.name)


def parse_payment_workbook(path: str | Path, *, source_reference: str | None = None) -> tuple[ParsedSocinproPayment, ...]:
    """Read the proven legacy 'pagamentos' projection without writing it."""
    candidate = Path(path)
    if candidate.suffix.casefold() not in {".xlsx", ".xlsm"} or not candidate.is_file():
        raise SocinproIngestionError("WORKBOOK_SOCINPRO_INVALIDO")
    try:
        book = load_workbook(candidate, read_only=True, data_only=False, keep_links=True)
        payment_sheet = next((name for name in book.sheetnames if _key(name) == "pagamentos"), None)
        if payment_sheet is None:
            raise SocinproIngestionError("ABA_PAGAMENTOS_AUSENTE")
        sheet = book[payment_sheet]
        header_row, headers = _find_payment_headers(sheet)
        aliases = {
            "source_code": ("cod_socinpro", "codigo na associacao"),
            "titular": ("titular",),
            "payment_date": ("data_pagamento", "data do pagamento"),
            "payment_value": ("valor_pagamento", "valor"),
            "document": ("arquivo_analitico", "documento"),
        }
        resolved = {name: next((headers[item] for item in options if item in headers), None) for name, options in aliases.items()}
        missing = [name for name in ("source_code", "titular", "payment_date", "payment_value") if resolved[name] is None]
        if missing:
            raise SocinproIngestionError("CAMPOS_WORKBOOK_AUSENTES:" + ",".join(missing))
        rows: list[ParsedSocinproPayment] = []
        for row_number, values in enumerate(sheet.iter_rows(min_row=header_row + 1, values_only=True), start=header_row + 1):
            if not any(value is not None and str(value).strip() for value in values):
                continue
            source_code = str(values[resolved["source_code"]] or "").strip()
            titular = str(values[resolved["titular"]] or "").strip()
            payment_date = _parse_date(values[resolved["payment_date"]])
            original = _parse_brl(values[resolved["payment_value"]])
            _require_payment_fields(source_code, titular, payment_date, original)
            document_index = resolved["document"]
            document = str(values[document_index] or "").strip() if document_index is not None else candidate.name
            reference = source_reference or f"{candidate.as_uri()}#pagamentos!{row_number}"
            rows.append(ParsedSocinproPayment(source_code, titular, payment_date, original, -original, document or candidate.name, reference))
        if not rows:
            raise SocinproIngestionError("WORKBOOK_SOCINPRO_SEM_PAGAMENTOS")
        return tuple(rows)
    finally:
        try: book.close()
        except UnboundLocalError: pass


def load_mapping_from_environment(environ: Mapping[str, str] | None = None) -> tuple[SocinproMapping, ...]:
    environment = os.environ if environ is None else environ
    raw_path = str(environment.get(MAPPING_ENVIRONMENT_VARIABLE) or "").strip()
    if not raw_path:
        raise SocinproIngestionError("MAPPING_SOCINPRO_NAO_CONFIGURADO")
    return load_mapping_file(raw_path)


def load_mapping_file(path: str | Path) -> tuple[SocinproMapping, ...]:
    candidate = Path(path)
    if not candidate.is_file():
        raise SocinproIngestionError("MAPPING_SOCINPRO_NAO_ENCONTRADO")
    try:
        payload = json.loads(candidate.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as exc:
        raise SocinproIngestionError("MAPPING_SOCINPRO_INVALIDO") from exc
    if payload.get("schema_version") != 1 or not isinstance(payload.get("mappings"), list):
        raise SocinproIngestionError("SCHEMA_MAPPING_SOCINPRO_INVALIDO")
    mappings: list[SocinproMapping] = []
    keys: set[tuple[str, str]] = set()
    for item in payload["mappings"]:
        if not isinstance(item, dict): raise SocinproIngestionError("ITEM_MAPPING_SOCINPRO_INVALIDO")
        entity, code = str(item.get("entity") or "").upper().strip(), str(item.get("source_code") or "").strip()
        titular, catalog, deal = (str(item.get(name) or "").strip() for name in ("titular", "catalog", "deal"))
        active = item.get("active")
        if entity not in {"HM", "MDB"} or not code or not titular or not catalog or not deal or not isinstance(active, bool):
            raise SocinproIngestionError("ITEM_MAPPING_SOCINPRO_INVALIDO")
        key = (entity, code)
        if key in keys: raise SocinproIngestionError("MAPPING_SOCINPRO_AMBIGUO")
        keys.add(key); mappings.append(SocinproMapping(entity, code, titular, catalog, deal, active))
    return tuple(mappings)


def mapping_resolver(mappings: Iterable[SocinproMapping], entity: str) -> Callable[[SocinproPayment], CatalogRelation | None]:
    index = {(item.entity, item.source_code): item for item in mappings}
    normalized_entity = entity.upper().strip()
    def resolve(payment: SocinproPayment) -> CatalogRelation | None:
        item = index.get((normalized_entity, payment.source_code))
        if item is None or not item.active or _key(item.titular) != _key(payment.titular): return None
        return CatalogRelation(item.catalog, item.deal)
    return resolve


def _parse_payment_text(text: str, reference: str, document: str) -> ParsedSocinproPayment:
    lines = [re.sub(r"\s+", " ", line).strip() for line in text.splitlines() if line.strip()]
    if not any("pagamento efetuado" in _key(line) for line in lines):
        raise SocinproIngestionError("MARCADOR_PAGAMENTO_AUSENTE")
    titular_index = next((index for index, line in enumerate(lines) if "demonstrativo do titular" in _key(line)), -1)
    titular = lines[titular_index + 1] if titular_index >= 0 and titular_index + 1 < len(lines) else ""
    code_match = re.search(r"C[ÓO]D\.?(?:\s+|\n)*SOCINPRO\s*:\s*(?:C[ÓO]D\.?(?:\s+|\n)*ECAD\s*:\s*)?(\d+)", text, re.IGNORECASE)
    source_code = code_match.group(1) if code_match else ""
    payment_match = re.search(r"(\d{2}/\d{2}/\d{4})\s+(-?\s*R\$\s*-?\s*\d{1,3}(?:\.\d{3})*,\d{2})\s*Pagamento\s+efetuado", text, re.IGNORECASE)
    if not payment_match:
        payment_match = re.search(r"Pagamento\s+efetuado\s*:\s*(\d{2}/\d{2}/\d{4})\s*-?\s*(-?\s*R\$\s*-?\s*\d{1,3}(?:\.\d{3})*,\d{2})", text, re.IGNORECASE)
    if not payment_match: raise SocinproIngestionError("DATA_OU_VALOR_PAGAMENTO_AUSENTE")
    payment_date, original = _parse_date(payment_match.group(1)), _parse_brl(payment_match.group(2))
    _require_payment_fields(source_code, titular, payment_date, original)
    return ParsedSocinproPayment(source_code, titular, payment_date, original, -original, document, reference)


def _find_payment_headers(sheet) -> tuple[int, dict[str, int]]:
    """Find the approved legacy projection or operational display header."""
    for row_number, values in enumerate(sheet.iter_rows(min_row=1, max_row=12, values_only=True), start=1):
        headers = {_key(value): index for index, value in enumerate(values) if value is not None}
        if "titular" in headers and ("cod_socinpro" in headers or "codigo na associacao" in headers):
            return row_number, headers
    raise SocinproIngestionError("CAMPOS_WORKBOOK_AUSENTES:cabecalho")


def _require_payment_fields(code: str, titular: str, payment_date: date, original: Decimal) -> None:
    if not code: raise SocinproIngestionError("CODIGO_SOCINPRO_AUSENTE")
    if not titular: raise SocinproIngestionError("TITULAR_SOCINPRO_AUSENTE")
    if original >= 0: raise SocinproIngestionError("SINAL_SOCINPRO_INESPERADO")


def _parse_date(value: object) -> date:
    if isinstance(value, datetime): return value.date()
    if isinstance(value, date): return value
    try: return datetime.strptime(str(value).strip(), "%d/%m/%Y").date()
    except ValueError:
        try: return date.fromisoformat(str(value).strip())
        except ValueError as exc: raise SocinproIngestionError("DATA_SOCINPRO_INVALIDA") from exc


def _parse_brl(value: object) -> Decimal:
    # openpyxl exposes numeric Excel cells as float. Convert its displayed
    # decimal representation immediately; no float reaches the typed contract.
    if isinstance(value, float):
        return Decimal(str(value))
    raw = str(value).strip().replace("R$", "").replace(" ", "")
    if not raw: raise SocinproIngestionError("VALOR_SOCINPRO_AUSENTE")
    negative = "-" in raw
    normalized = raw.replace("-", "").replace(".", "").replace(",", ".")
    try: return -Decimal(normalized) if negative else Decimal(normalized)
    except InvalidOperation as exc: raise SocinproIngestionError("VALOR_SOCINPRO_INVALIDO") from exc


def _semantic_identity(code: str, titular: str, payment_date: date, original: Decimal) -> str:
    raw = "|".join((code, _key(titular), payment_date.isoformat(), format(original, "f")))
    return sha256(raw.encode("utf-8")).hexdigest()


def _key(value: object) -> str:
    raw = unicodedata.normalize("NFKD", str(value or ""))
    return re.sub(r"\s+", " ", "".join(char for char in raw if not unicodedata.combining(char))).strip().casefold()
