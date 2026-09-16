"""Source-neutral demonstrativo service; parsing and publication stay outside it."""
from __future__ import annotations

from dataclasses import dataclass
from datetime import date
from decimal import Decimal
from pathlib import Path
from typing import Iterable

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill


@dataclass(frozen=True)
class DemonstrativoRow:
    payment_date: date
    source_code: str
    titular: str
    canonical_artist: str
    value: Decimal
    document: str
    status: str
    observation: str = ""


def generate_demonstrativo(*, entity: str, source: str, competence: str, financial_rows: Iterable[DemonstrativoRow], review_items: Iterable[str], provenance: Iterable[str], destination: str | Path) -> Path:
    """Write a review-only XLSX from typed financial rows.

    It deliberately has no SOCINPRO parser, bank connector, or publication
    behavior, so callers for HM/MDB and future associations share one contract.
    """
    entity, source = entity.upper().strip(), source.upper().strip()
    if entity not in {"HM", "MDB"} or not source or len(competence) != 7:
        raise ValueError("DEMONSTRATIVO_CONTEXT_INVALID")
    rows = tuple(financial_rows)
    if any(row.value <= Decimal("0") for row in rows):
        raise ValueError("DEMONSTRATIVO_VALUE_INVALID")
    target = Path(destination)
    target.parent.mkdir(parents=True, exist_ok=True)
    book = Workbook(); payments = book.active; payments.title = "Pagamentos"
    headers = ["Data pagamento", "Código fonte/associação", "Titular", "Artista / Titular canônico", "Valor", "Documento", "Status", "Observação"]
    payments.append(headers)
    for row in rows:
        payments.append([row.payment_date, row.source_code, row.titular, row.canonical_artist, row.value, row.document, row.status, row.observation])
    summary = book.create_sheet("Resumo por Titular")
    summary.append(["Artista / Titular canônico", "Quantidade", "Valor"])
    aggregates: dict[str, tuple[int, Decimal]] = {}
    for row in rows:
        label = row.canonical_artist or row.titular
        count, value = aggregates.get(label, (0, Decimal("0.00")))
        aggregates[label] = (count + 1, value + row.value)
    for label in sorted(aggregates, key=str.casefold):
        count, value = aggregates[label]; summary.append([label, count, value])
    pending = book.create_sheet("Pendências")
    pending.append(["Pendência"])
    for item in review_items:
        pending.append([item])
    fill = PatternFill("solid", fgColor="1F4E78"); font = Font(color="FFFFFF", bold=True)
    for sheet in book.worksheets:
        for cell in sheet[1]: cell.fill = fill; cell.font = font
        sheet.freeze_panes = "A2"
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = min(45, max(12, max(len(str(cell.value or "")) for cell in column) + 2))
    for cell in payments["A"][1:]: cell.number_format = "DD/MM/YYYY"
    for cell in payments["E"][1:]: cell.number_format = '#,##0.00'
    for cell in summary["C"][1:]: cell.number_format = '#,##0.00'
    book.save(target); book.close()
    return target
