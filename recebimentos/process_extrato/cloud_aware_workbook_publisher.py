"""Cloud-aware, fail-closed publication for official XLSX workbooks.

This module never treats a synced/reparse-point path as safely replaceable by
local overwrite.  Cloud-native transport is deliberately an injected future
adapter: without it, publication is BLOCKED.
"""
from __future__ import annotations

from dataclasses import asdict, dataclass
from enum import Enum
from hashlib import sha256
import ctypes
import json
import os
from pathlib import Path
import shutil
import uuid
from zipfile import ZipFile

from openpyxl import load_workbook


class StorageClass(str, Enum):
    LOCAL_NORMAL = "LOCAL_NORMAL"
    SYNCED_LOCAL_ATOMIC_CAPABLE = "SYNCED_LOCAL_ATOMIC_CAPABLE"
    SYNCED_CLOUD_REPARSE = "SYNCED_CLOUD_REPARSE"
    BLOCKED = "BLOCKED"


@dataclass(frozen=True)
class StorageDiagnosis:
    path: str
    storage_provider: str
    storage_class: StorageClass
    is_reparse_point: bool
    is_locally_hydrated: bool | None
    locked: bool
    write_permission: bool
    can_atomically_replace_existing_file: bool | None = None
    errors: tuple[str, ...] = ()


@dataclass(frozen=True)
class DirectoryCapability:
    directory: str
    binary_copy: bool
    rename: bool
    replace_existing: bool
    rename_replace_existing: bool
    can_atomically_replace_existing_file: bool
    errors: tuple[str, ...] = ()


@dataclass(frozen=True)
class PublicationResult:
    publication_method: str
    status: str
    source_hash: str | None
    destination_before: str | None
    destination_after: str | None
    semantic_fingerprint: str | None
    rollback_available: bool
    warnings: tuple[str, ...] = ()
    errors: tuple[str, ...] = ()
    backup_path: str | None = None
    storage_type: str | None = None


FILE_ATTRIBUTE_REPARSE_POINT = 0x0400
INVALID_FILE_ATTRIBUTES = 0xFFFFFFFF


def binary_hash(path: str | Path) -> str:
    digest = sha256()
    with Path(path).open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _win32_attributes(path: Path) -> int | None:
    # Python exposes the authoritative Windows attributes without losing a
    # Unicode path to a ctypes conversion on synced folders.
    try:
        attributes = getattr(path.stat(), "st_file_attributes", None)
        if attributes is not None:
            return int(attributes)
    except OSError:
        pass
    if os.name != "nt":
        return None
    function = ctypes.windll.kernel32.GetFileAttributesW
    function.argtypes = (ctypes.c_wchar_p,)
    function.restype = ctypes.c_uint32
    value = function(str(path))
    return None if value == INVALID_FILE_ATTRIBUTES else int(value)


def _provider(path: Path, reparse: bool) -> str:
    folded = str(path).casefold()
    if "onedrive" in folded:
        return "OneDrive"
    if "sharepoint" in folded or " - documentos" in folded:
        return "SharePoint/OneDrive synced"
    return "Unknown synced provider" if reparse else "Local filesystem"


def _is_locked(path: Path) -> bool:
    if not path.exists():
        return False
    try:
        with path.open("r+b"):
            return False
    except PermissionError:
        return True


def _can_write(directory: Path) -> bool:
    probe = directory / f".muv-write-probe-{uuid.uuid4().hex}.tmp"
    try:
        probe.write_bytes(b"probe")
        probe.unlink()
        return True
    except OSError:
        return False


def diagnose_storage(path: str | Path) -> StorageDiagnosis:
    """Read-only target diagnosis plus a reversible directory write probe."""
    target = Path(path)
    attributes = _win32_attributes(target)
    reparse = bool(attributes is not None and attributes & FILE_ATTRIBUTE_REPARSE_POINT)
    provider = _provider(target, reparse)
    locked = _is_locked(target)
    writable = _can_write(target.parent)
    errors: list[str] = []
    if not target.exists():
        errors.append("DESTINATION_NOT_FOUND")
    if locked:
        errors.append("DESTINATION_LOCKED")
    if not writable:
        errors.append("DIRECTORY_NOT_WRITABLE")
    # A synced provider name is itself enough to deny local publication.  Some
    # embedded Python runtimes mask FILE_ATTRIBUTE_REPARSE_POINT, so absence of
    # that bit must never downgrade a known OneDrive/SharePoint path to local.
    synced_path = provider != "Local filesystem"
    storage_class = StorageClass.SYNCED_CLOUD_REPARSE if reparse or synced_path else StorageClass.LOCAL_NORMAL
    if locked or not writable:
        storage_class = StorageClass.BLOCKED
    return StorageDiagnosis(str(target), provider, storage_class, reparse, None, locked, writable, errors=tuple(errors))


def _error(label: str, exc: OSError) -> str:
    return f"{label}: winerror={getattr(exc, 'winerror', None)} errno={exc.errno} {exc}"


def probe_directory_capability(directory: str | Path) -> DirectoryCapability:
    """Prove replace-existing behavior using only fresh sibling dummy files."""
    folder = Path(directory)
    token = uuid.uuid4().hex
    source, renamed, target, replacement = (folder / f".muv-publish-probe-{token}-{name}.tmp" for name in ("source", "renamed", "target", "replacement"))
    results = {"binary_copy": False, "rename": False, "replace_existing": False, "rename_replace_existing": False}
    errors: list[str] = []
    try:
        source.write_bytes(b"source")
        shutil.copyfile(source, target); results["binary_copy"] = target.read_bytes() == b"source"
        source.replace(renamed); results["rename"] = renamed.read_bytes() == b"source"
        target.write_bytes(b"target")
        replacement.write_bytes(b"replacement")
        try:
            os.replace(replacement, target)
            results["replace_existing"] = target.read_bytes() == b"replacement"
        except OSError as exc:
            errors.append(_error("REPLACE_EXISTING", exc))
        replacement.write_bytes(b"rename-replacement")
        try:
            replacement.replace(target)
            results["rename_replace_existing"] = target.read_bytes() == b"rename-replacement"
        except OSError as exc:
            errors.append(_error("RENAME_REPLACE_EXISTING", exc))
    except OSError as exc:
        errors.append(_error("PROBE", exc))
    finally:
        for file in (source, renamed, target, replacement):
            try: file.unlink(missing_ok=True)
            except OSError as exc: errors.append(_error("CLEANUP", exc))
    capable = results["replace_existing"] and results["rename_replace_existing"]
    return DirectoryCapability(str(folder), **results, can_atomically_replace_existing_file=capable, errors=tuple(errors))


def semantic_fingerprint(path: str | Path) -> str:
    """Stable workbook-content fingerprint that ignores OOXML package metadata."""
    workbook = load_workbook(Path(path), read_only=True, data_only=False, keep_links=True)
    content: list[dict[str, object]] = []
    try:
        for sheet in workbook.worksheets:
            cells = []
            for row in sheet.iter_rows():
                for cell in row:
                    if cell.value is not None:
                        cells.append((cell.coordinate, str(cell.value), cell.data_type))
            content.append({"title": sheet.title, "state": sheet.sheet_state, "cells": cells})
    finally:
        workbook.close()
    return sha256(json.dumps(content, ensure_ascii=False, separators=(",", ":")).encode("utf-8")).hexdigest()


def classify_xlsx_diff(before: str | Path, after: str | Path) -> str:
    """Metadata-only is valid only when the semantic fingerprint agrees."""
    if semantic_fingerprint(before) != semantic_fingerprint(after):
        return "SEMANTIC_CHANGE"
    with ZipFile(before) as left, ZipFile(after) as right:
        changed = set(left.namelist()) ^ set(right.namelist())
        for name in set(left.namelist()) & set(right.namelist()):
            if left.read(name) != right.read(name): changed.add(name)
    metadata = ("docProps/", "customXml/", "_rels/", "[Content_Types].xml")
    return "OFFICE_METADATA_ONLY" if all(item.startswith(metadata) for item in changed) else "BINARY_ONLY_SEMANTIC_MATCH"


def publish_workbook(source_path: str | Path, destination_path: str | Path, expected_source_hash: str, expected_destination_fingerprint: str | None = None) -> PublicationResult:
    """Publish locally only after capability and concurrency guards pass.

    Cloud/reparse destinations are fail-closed until an authenticated remote
    publisher with ETag/version guards is supplied by the caller.
    """
    source, destination = Path(source_path), Path(destination_path)
    source_hash = binary_hash(source)
    diagnosis = diagnose_storage(destination)
    if source_hash != expected_source_hash:
        return PublicationResult("BLOCKED", "BLOCKED", source_hash, None, None, None, False, errors=("SOURCE_HASH_MISMATCH",), storage_type=diagnosis.storage_class.value)
    if diagnosis.storage_class is not StorageClass.LOCAL_NORMAL:
        return PublicationResult("BLOCKED", "BLOCKED", source_hash, None, None, None, False, warnings=("LOCAL_ATOMIC_PUBLICATION_UNAVAILABLE",), errors=diagnosis.errors or ("CLOUD_NATIVE_PUBLISHER_UNAVAILABLE",), storage_type=diagnosis.storage_class.value)
    before_fingerprint = semantic_fingerprint(destination)
    if expected_destination_fingerprint is None or before_fingerprint != expected_destination_fingerprint:
        return PublicationResult("BLOCKED", "BLOCKED", source_hash, before_fingerprint, None, None, False, errors=("CONCURRENCY_GUARD_FAILED",), storage_type=diagnosis.storage_class.value)
    capability = probe_directory_capability(destination.parent)
    if not capability.can_atomically_replace_existing_file:
        return PublicationResult("BLOCKED", "BLOCKED", source_hash, before_fingerprint, None, None, False, errors=capability.errors or ("LOCAL_ATOMIC_PUBLICATION_UNAVAILABLE",), storage_type=diagnosis.storage_class.value)
    backup = destination.with_name(destination.name + f".muv-backup-{uuid.uuid4().hex}")
    sibling = destination.with_name(destination.name + f".muv-staging-{uuid.uuid4().hex}")
    try:
        shutil.copy2(destination, backup)
        shutil.copy2(source, sibling)
        if binary_hash(sibling) != source_hash: raise RuntimeError("TEMP_HASH_MISMATCH")
        os.replace(sibling, destination)
        after_fingerprint = semantic_fingerprint(destination)
        source_fingerprint = semantic_fingerprint(source)
        if after_fingerprint != source_fingerprint: raise RuntimeError("POST_PUBLICATION_SEMANTIC_MISMATCH")
        return PublicationResult("LOCAL_ATOMIC", "PASS", source_hash, before_fingerprint, binary_hash(destination), after_fingerprint, True, backup_path=str(backup), storage_type=diagnosis.storage_class.value)
    except Exception as exc:
        try:
            if backup.exists(): os.replace(backup, destination)
        except OSError as rollback_error:
            return PublicationResult("LOCAL_ATOMIC", "FAIL", source_hash, before_fingerprint, None, None, False, errors=(str(exc), _error("ROLLBACK", rollback_error)), storage_type=diagnosis.storage_class.value)
        return PublicationResult("LOCAL_ATOMIC", "FAIL", source_hash, before_fingerprint, None, None, True, errors=(str(exc),), backup_path=str(backup), storage_type=diagnosis.storage_class.value)
    finally:
        sibling.unlink(missing_ok=True)


def create_workbook_only(source_path: str | Path, destination_path: str | Path) -> PublicationResult:
    """Create a new workbook without any overwrite path.

    This shares the publisher's synced-folder policy.  The exclusive create
    itself is the concurrency guard: if another process creates the month
    between discovery and publication, this returns ``CONFLICT`` unchanged.
    """
    source, destination = Path(source_path), Path(destination_path)
    if not source.is_file():
        return PublicationResult("CREATE_ONLY", "BLOCKED", None, None, None, None, False, errors=("SOURCE_NOT_FOUND",))
    if destination.exists():
        return PublicationResult("CREATE_ONLY", "CONFLICT", binary_hash(source), str(destination), None, None, False, errors=("DESTINATION_ALREADY_EXISTS",))
    destination.parent.mkdir(parents=True, exist_ok=True)
    diagnosis = diagnose_storage(destination)
    if diagnosis.storage_class is not StorageClass.LOCAL_NORMAL:
        return PublicationResult("CREATE_ONLY", "BLOCKED", binary_hash(source), None, None, None, False, errors=diagnosis.errors or ("CREATE_ONLY_UNAVAILABLE",), storage_type=diagnosis.storage_class.value)
    source_hash = binary_hash(source)
    try:
        handle = os.open(destination, os.O_WRONLY | os.O_CREAT | os.O_EXCL)
    except FileExistsError:
        return PublicationResult("CREATE_ONLY", "CONFLICT", source_hash, str(destination), None, None, False, errors=("DESTINATION_ALREADY_EXISTS",), storage_type=diagnosis.storage_class.value)
    try:
        with os.fdopen(handle, "wb") as stream:
            stream.write(source.read_bytes())
        if binary_hash(destination) != source_hash:
            destination.unlink(missing_ok=True)
            return PublicationResult("CREATE_ONLY", "BLOCKED", source_hash, None, None, None, False, errors=("POST_CREATE_HASH_MISMATCH",), storage_type=diagnosis.storage_class.value)
        return PublicationResult("CREATE_ONLY", "PASS", source_hash, None, str(destination), semantic_fingerprint(destination), False, storage_type=diagnosis.storage_class.value)
    except Exception as exc:
        destination.unlink(missing_ok=True)
        return PublicationResult("CREATE_ONLY", "BLOCKED", source_hash, None, None, None, False, errors=(str(exc),), storage_type=diagnosis.storage_class.value)


def checkpoint(path: str | Path) -> dict[str, object]:
    diagnosis = diagnose_storage(path)
    capability = probe_directory_capability(Path(path).parent) if diagnosis.write_permission else None
    local_atomic = diagnosis.storage_class is StorageClass.LOCAL_NORMAL and bool(capability and capability.can_atomically_replace_existing_file)
    return {"CLOUD_AWARE_WORKBOOK_PUBLISHER_STATUS": "PASS", "STORAGE_PROVIDER": diagnosis.storage_provider, "IS_REPARSE_POINT": diagnosis.is_reparse_point, "IS_LOCALLY_HYDRATED": diagnosis.is_locally_hydrated, "CAN_ATOMICALLY_REPLACE_EXISTING_FILE": capability.can_atomically_replace_existing_file if capability else False, "LOCAL_ATOMIC_PUBLICATION_AVAILABLE": local_atomic, "CLOUD_NATIVE_PUBLISHER_AVAILABLE": False, "SELECTED_PUBLICATION_METHOD": "LOCAL_ATOMIC" if local_atomic else "BLOCKED", "NON_ATOMIC_OVERWRITE_USED": False, "CONCURRENCY_GUARD_AVAILABLE": True, "ROLLBACK_AVAILABLE": local_atomic, "SYNC_AWARE_VALIDATION_AVAILABLE": True, "SEMANTIC_FINGERPRINT_SUPPORTED": True, "OFFICIAL_FILES_CHANGED_DURING_DIAGNOSIS": False, "diagnosis": asdict(diagnosis), "probe_errors": capability.errors if capability else ()}
